from dotenv import load_dotenv
load_dotenv()

from flask import Flask, render_template, request, redirect, url_for, send_file, session, flash
from werkzeug.security import generate_password_hash, check_password_hash
from itsdangerous import URLSafeTimedSerializer, BadSignature, SignatureExpired
from urllib.parse import quote
from datetime import datetime, timedelta

# Conexão: produção = Turso; local = SQL Server (pyodbc) ou SQLite
import db_adapter
import auth_service as auth
import agenda_service as agenda
from database import (
    DbError,
    SERVICOS_PADRAO,
    _coluna_existe,
    _valor_linha,
    ensure_database_schema,
    ensure_schema_migrations,
    ensure_schema_migrations_conn,
    garantir_banco_pronto,
    garantir_comissoes_defaults,
    get_connection,
    init_database,
    initialize_database_schema,
    obter_id_inserido,
    safe_close,
    safe_commit,
    safe_rollback,
    seed_servicos_horarios_padrao,
    vincular_usuario_barbearia,
)
import smtplib
from email.mime.text import MIMEText
from email.mime.multipart import MIMEMultipart
from subscriptions import criar_assinatura_trial, requer_assinatura_ativa
from stripe_payments import register_stripe_routes
import config_saas as cfg
import financeiro_service as fin
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas
from reportlab.lib.units import cm
import socket
import io
import re
import os
import secrets
import traceback

# 1. IMPORTAR O FLASK-BABEL
from flask_babel import Babel, _, gettext

app = Flask(__name__)
app.secret_key = (
    os.environ.get("SECRET_KEY")
    or os.environ.get("FLASK_SECRET_KEY")
    or "dev-only-change-me-in-production"
)

# 2. CONFIGURAR O FLASK-BABEL (PT-BR, EN-US, ES-ES)
app.config["BABEL_DEFAULT_LOCALE"] = "pt"
app.config["BABEL_SUPPORTED_LOCALES"] = ["pt", "en", "es"]
app.config["BABEL_TRANSLATION_DIRECTORIES"] = "translations"


def get_locale():
    """Prioridade: sessão do usuário → idioma do navegador → português."""
    lang = session.get("lang")
    if lang in app.config["BABEL_SUPPORTED_LOCALES"]:
        return lang
    match = request.accept_languages.best_match(app.config["BABEL_SUPPORTED_LOCALES"])
    return match or app.config["BABEL_DEFAULT_LOCALE"]


babel = Babel(app, locale_selector=get_locale)


@app.context_processor
def inject_i18n():
    ctx = {
        "current_locale": get_locale(),
        "supported_locales": app.config["BABEL_SUPPORTED_LOCALES"],
    }
    if session.get("role") == "admin" and session.get("barbearia_id"):
        try:
            from subscriptions import resumo_plano_admin

            ctx["plano_resumo"] = resumo_plano_admin(session["barbearia_id"])
        except Exception:
            ctx["plano_resumo"] = {"mostrar_banner": False}
    return ctx

HORARIOS = [
    "08:00", "09:00", "10:00", "11:00",
    "12:00", "13:00", "14:00", "15:00",
    "16:00", "17:00", "18:00", "19:00",
    "20:00"
]

DIAS_PT = {
    "Monday": "Segunda-feira",
    "Tuesday": "Terça-feira",
    "Wednesday": "Quarta-feira",
    "Thursday": "Quinta-feira",
    "Friday": "Sexta-feira",
    "Saturday": "Sábado",
    "Sunday": "Domingo"
}

# -------------------------- RECUPERAÇÃO DE SENHA --------------------------
_RESET_TOKEN_MAX_AGE = 30 * 60  # 30 minutos
_reset_serializer = None


def _get_reset_serializer():
    global _reset_serializer
    if _reset_serializer is None:
        _reset_serializer = URLSafeTimedSerializer(
            app.secret_key, salt="password-reset-v1"
        )
    return _reset_serializer


def _gerar_token_recuperacao(user_id):
    return _get_reset_serializer().dumps({"uid": int(user_id)})


def _validar_token_recuperacao(token):
    try:
        data = _get_reset_serializer().loads(token, max_age=_RESET_TOKEN_MAX_AGE)
        return int(data["uid"])
    except (BadSignature, SignatureExpired, KeyError, TypeError, ValueError):
        return None


def _limpar_telefone(valor):
    return re.sub(r"\D", "", (valor or "").strip())


PERFIL_UPLOAD_DIR = "static/uploads/perfil"
_EXTENSOES_FOTO_PERFIL = frozenset({"jpg", "jpeg", "png", "webp", "gif"})


def _open_db():
    """Abre conexão com schema garantido e cursor que traduz SQL (Turso/SQL Server)."""
    conn = db_adapter.get_connection()
    db_adapter.ensure_app_schema(conn)
    return conn, db_adapter.cursor(conn)


def _coluna_usuarios_existe(cursor, coluna):
    conn = getattr(cursor, "_conn", None)
    return db_adapter.coluna_existe("usuarios", coluna, conn=conn)


def _coluna_existe(cursor, tabela, coluna):
    conn = getattr(cursor, "_conn", None)
    return db_adapter.coluna_existe(tabela, coluna, conn=conn)


def _extensao_imagem_segura(filename):
    if not filename or "." not in filename:
        return None
    ext = filename.rsplit(".", 1)[1].lower()
    if ext not in _EXTENSOES_FOTO_PERFIL:
        return None
    return ext


def _salvar_upload_foto_perfil(arquivo):
    """Salva foto em static/uploads/perfil; retorna nome do arquivo ou None."""
    if not arquivo or not arquivo.filename:
        return None
    ext = _extensao_imagem_segura(arquivo.filename)
    if not ext:
        return None
    os.makedirs(PERFIL_UPLOAD_DIR, exist_ok=True)
    nome = f"perfil_{secrets.token_hex(12)}.{ext}"
    arquivo.save(os.path.join(PERFIL_UPLOAD_DIR, nome))
    return nome


ROLES_EQUIPE = frozenset({"admin", "barbeiro", "profissional"})
MSG_ERRO_VINCULO = "Erro de vinculação: contate o administrador"


def _definir_sessao_usuario(
    user_id, user_name, role, barbearia_id, nome_barbearia, barbearia_slug
):
    """Grava sessão Flask com tipos JSON-safe (role real: admin, barbeiro, profissional)."""
    role_norm = (str(role or "profissional").strip().lower())
    if role_norm not in ROLES_EQUIPE:
        role_norm = "profissional"

    session.clear()
    session["user_id"] = int(user_id)
    session["usuario_id"] = int(user_id)
    session["user_name"] = str(user_name or "").strip() or "Usuário"
    session["role"] = role_norm
    session["tipo_usuario"] = role_norm
    session["barbearia_id"] = int(barbearia_id)
    session["nome_barbearia"] = str(nome_barbearia or "").strip() or "AgendaSimples"
    session["barbearia_slug"] = str(barbearia_slug or "").strip()
    session.modified = True


def _definir_sessao_admin(user_id, user_name, barbearia_id, nome_barbearia, barbearia_slug):
    """Atalho para sessão de administrador do negócio."""
    _definir_sessao_usuario(
        user_id, user_name, "admin", barbearia_id, nome_barbearia, barbearia_slug
    )


def _role_sessao():
    return (session.get("role") or "").strip().lower()


def _usuario_autenticado():
    return bool(session.get("user_id")) and bool(session.get("barbearia_id"))


def _eh_admin():
    return _role_sessao() == "admin"


def _eh_profissional_equipe():
    return _role_sessao() in ("barbeiro", "profissional")


def _barbearia_id_sessao():
    """ID do negócio na sessão (admin ou profissional)."""
    return session.get("barbearia_id")


def _exigir_login_barbearia():
    """Usuário logado com barbearia vinculada — retorna redirect ou None."""
    if not session.get("user_id"):
        flash(_("Faça login para continuar."), "warning")
        return redirect(url_for("login"))
    if not session.get("barbearia_id"):
        flash(_("Sessão inválida. Faça login novamente."), "error")
        return redirect(url_for("login"))
    if _role_sessao() not in ROLES_EQUIPE:
        flash(_("Você não tem permissão para acessar esta área."), "error")
        return redirect(url_for("acesso_negado"))
    return None


def _exigir_admin():
    """Somente administrador do negócio."""
    bloqueio = _exigir_login_barbearia()
    if bloqueio:
        return bloqueio
    if not _eh_admin():
        flash(_("Esta página é exclusiva do administrador."), "warning")
        return redirect(url_for("acesso_negado"))
    return None


def _linha_usuario_login(row):
    """
    Extrai campos do SELECT de login (id, nome, role, senha, barbearia_id).
    Usa índice e nome de coluna — compatível com Turso HTTP e SQLite.
    """
    if row is None:
        return None
    uid = _valor_linha(row, 0, "id")
    nome = _valor_linha(row, 1, "nome")
    role = _valor_linha(row, 2, "role") or _valor_linha(row, nome="role")
    senha = _valor_linha(row, 3, "senha") or _valor_linha(row, nome="senha")
    barbearia_id = _valor_linha(row, 4, "barbearia_id") or _valor_linha(
        row, nome="barbearia_id"
    )
    return {
        "id": uid,
        "nome": nome,
        "role": role,
        "senha": senha,
        "barbearia_id": barbearia_id,
    }


def _barbearia_id_do_usuario(cursor, user_id, tentar_corrigir=True):
    """barbearia_id gravado no registro do usuário (fonte de verdade)."""
    conn = getattr(cursor, "_conn", None)
    if conn is not None:
        return auth.barbearia_id_do_usuario(
            int(user_id), conn=conn, tentar_corrigir=tentar_corrigir
        )
    return auth.barbearia_id_do_usuario(int(user_id), tentar_corrigir=tentar_corrigir)


def _logout_por_falta_vinculo():
    """Encerra sessão e impede acesso sem barbearia_id no banco."""
    session.clear()
    flash(_(MSG_ERRO_VINCULO), "error")
    return redirect(url_for("login"))


def _barbearia_id_admin_sessao_obrigatorio():
    """barbearia_id do admin logado — None se ausente ou inválido."""
    bid = session.get("barbearia_id")
    if bid is None:
        return None
    try:
        return int(bid)
    except (TypeError, ValueError):
        return None


def _sincronizar_barbearia_sessao(cursor):
    """
    Garante que session['barbearia_id'] coincide com usuarios.barbearia_id.
    Retorna redirect se o usuário continuar órfão.
    """
    user_id = session.get("user_id") or session.get("usuario_id")
    if not user_id:
        return redirect(url_for("login"))

    bid_db = _barbearia_id_do_usuario(cursor, user_id, tentar_corrigir=False)
    if bid_db is None:
        return _logout_por_falta_vinculo()

    bid_sessao = session.get("barbearia_id")
    if bid_sessao is None or int(bid_sessao) != int(bid_db):
        session["barbearia_id"] = int(bid_db)
        session.modified = True
        if bid_sessao is not None and int(bid_sessao) != int(bid_db):
            app.logger.warning(
                "Sessão corrigida: user_id=%s barbearia_id %s -> %s",
                user_id,
                bid_sessao,
                bid_db,
            )
    return None


def _inserir_profissional_usuario(
    cursor, nome, email, senha_hash, barbearia_id, telefone=None, especialidade=None, foto_perfil=None
):
    """INSERT em usuarios sempre com barbearia_id da sessão do admin."""
    if barbearia_id is None:
        raise ValueError("barbearia_id obrigatório para cadastrar profissional")
    try:
        bid = int(barbearia_id)
    except (TypeError, ValueError) as exc:
        raise ValueError("barbearia_id inválido") from exc
    if bid < 1:
        raise ValueError("barbearia_id inválido")

    cols = ["nome", "email", "senha", "role", "barbearia_id"]
    vals = [nome, email, senha_hash, "profissional", bid]
    if telefone is not None and _coluna_existe(cursor, "usuarios", "telefone"):
        cols.append("telefone")
        vals.append(telefone)
    if especialidade is not None and _coluna_existe(cursor, "usuarios", "especialidade"):
        cols.append("especialidade")
        vals.append(especialidade)
    if foto_perfil is not None and _coluna_existe(cursor, "usuarios", "foto_perfil"):
        cols.append("foto_perfil")
        vals.append(foto_perfil)
    placeholders = ", ".join("?" for _ in vals)
    colunas_sql = ", ".join(cols)
    cursor.execute(
        f"INSERT INTO usuarios ({colunas_sql}) VALUES ({placeholders})",
        tuple(vals),
    )


def _iniciar_sessao_usuario(cursor, user, email):
    """Preenche session Flask após login ou cadastro bem-sucedido."""
    if isinstance(user, dict):
        user_id = user.get("id")
        user_name = user.get("nome")
        role = user.get("role")
        barbearia_id = user.get("barbearia_id")
    else:
        parsed = _linha_usuario_login(user)
        if parsed:
            user_id = parsed["id"]
            user_name = parsed["nome"]
            role = parsed["role"]
            barbearia_id = parsed["barbearia_id"]
        else:
            user_id = _valor_linha(user, 0, "id")
            user_name = _valor_linha(user, 1, "nome")
            role = _valor_linha(user, nome="role")
            barbearia_id = _valor_linha(user, 4, "barbearia_id") or _valor_linha(
                user, nome="barbearia_id"
            )
    if role is None:
        if isinstance(user, dict):
            role = user.get("role", "profissional")
        else:
            try:
                role = user["role"]
            except (KeyError, TypeError, IndexError):
                role = getattr(user, "role", "profissional")
    role_norm = (str(role or "profissional").strip().lower())

    if barbearia_id is None and isinstance(user, dict):
        barbearia_id = user.get("barbearia_id")
    elif barbearia_id is None:
        try:
            barbearia_id = user["barbearia_id"]
        except (KeyError, TypeError, IndexError):
            barbearia_id = getattr(user, "barbearia_id", None)

    if barbearia_id is None and user_id:
        barbearia_id = vincular_usuario_barbearia(cursor, int(user_id))
        if barbearia_id is None:
            cursor.execute(
                "SELECT barbearia_id FROM usuarios WHERE id = ?",
                (int(user_id),),
            )
            ref = cursor.fetchone()
            if ref:
                barbearia_id = _valor_linha(ref, 0, "barbearia_id")

    b_row = None
    if barbearia_id:
        cursor.execute(
            "SELECT id, nome, slug FROM barbearias WHERE id = ?",
            (barbearia_id,),
        )
        b_row = cursor.fetchone()
    # Só admin pode inferir barbearia pelo e-mail da tabela barbearias (evita vazamento entre tenants)
    if not b_row and email and role_norm == "admin":
        cursor.execute(
            """
            SELECT id, nome, slug FROM barbearias
            WHERE LOWER(TRIM(email)) = LOWER(?) LIMIT 1
            """,
            (email,),
        )
        b_row = cursor.fetchone()
        if b_row and barbearia_id is None:
            barbearia_id = _valor_linha(b_row, 0, "id")
            cursor.execute(
                "UPDATE usuarios SET barbearia_id = ? WHERE id = ?",
                (barbearia_id, int(user_id)),
            )

    if b_row:
        _definir_sessao_usuario(
            user_id,
            user_name,
            role_norm,
            _valor_linha(b_row, 0, "id"),
            _valor_linha(b_row, 1, "nome"),
            _valor_linha(b_row, 2, "slug"),
        )
    elif barbearia_id:
        _definir_sessao_usuario(
            user_id, user_name, role_norm, barbearia_id, "AgendaSimples", ""
        )
    else:
        raise ValueError("Não foi possível vincular o usuário a um estabelecimento na sessão.")


def _email_ja_cadastrado(cursor, email, ignorar_id=None):
    if ignorar_id:
        cursor.execute(
            "SELECT id FROM usuarios WHERE LOWER(TRIM(email)) = LOWER(?) AND id <> ?",
            (email, ignorar_id),
        )
    else:
        cursor.execute(
            "SELECT id FROM usuarios WHERE LOWER(TRIM(email)) = LOWER(?)",
            (email,),
        )
    return cursor.fetchone() is not None


def _listar_profissionais(cursor, barbearia_id):
    cursor.execute(
        """
        SELECT id, nome FROM usuarios
        WHERE barbearia_id = ? AND role IN ('barbeiro', 'profissional')
        ORDER BY nome
        """,
        (barbearia_id,),
    )
    return cursor.fetchall()


def _listar_servicos(cursor, barbearia_id):
    """Lista nomes dos serviços (compatível com templates antigos)."""
    return [s["nome"] for s in _listar_servicos_detalhados(cursor, barbearia_id)]


def _listar_servicos_detalhados(cursor, barbearia_id):
    """Serviços com preço para agendamento e financeiro."""
    tem_preco = _coluna_existe(cursor, "servicos", "preco")
    if tem_preco:
        cursor.execute(
            """
            SELECT nome, IFNULL(preco, 0) FROM servicos
            WHERE barbearia_id = ? AND IFNULL(ativo, 1) = 1
            ORDER BY ordem, nome
            """,
            (barbearia_id,),
        )
        rows = cursor.fetchall()
        if rows:
            return [{"nome": r[0], "preco": float(r[1] or 0)} for r in rows]
    else:
        cursor.execute(
            """
            SELECT nome FROM servicos
            WHERE barbearia_id = ? AND IFNULL(ativo, 1) = 1
            ORDER BY ordem, nome
            """,
            (barbearia_id,),
        )
        rows = cursor.fetchall()
        if rows:
            return [
                {"nome": r[0], "preco": _preco_padrao_servico(r[0])} for r in rows
            ]
    return [
        {"nome": n, "preco": _preco_padrao_servico(n)} for n in SERVICOS_PADRAO
    ]


def _preco_padrao_servico(nome_servico):
    """Preço sugerido quando a coluna servicos.preco não existe ou é zero."""
    padroes = {
        "Corte": 50.0,
        "Corte e Barba": 70.0,
        "Corte + Sobrancelha": 55.0,
        "Barba": 35.0,
        "Outro": 0.0,
    }
    return float(padroes.get((nome_servico or "").strip(), 0.0))


def _parse_valor_monetario(valor_raw):
    if valor_raw is None or str(valor_raw).strip() == "":
        return None
    try:
        return float(str(valor_raw).strip().replace(",", "."))
    except (TypeError, ValueError):
        return None


def _obter_preco_servico(cursor, barbearia_id, nome_servico):
    nome = (nome_servico or "").strip()
    if not nome:
        return 0.0
    if _coluna_existe(cursor, "servicos", "preco"):
        cursor.execute(
            """
            SELECT IFNULL(preco, 0) FROM servicos
            WHERE barbearia_id = ? AND nome = ? AND IFNULL(ativo, 1) = 1
            LIMIT 1
            """,
            (barbearia_id, nome),
        )
        row = cursor.fetchone()
        if row and float(row[0] or 0) > 0:
            return float(row[0])
    return _preco_padrao_servico(nome)


def _descricao_financeiro_agendamento(servico, nome_cliente):
    servico = (servico or "Atendimento").strip()
    nome = (nome_cliente or "Cliente").strip()
    return f"Agendamento: {servico} | Cliente: {nome}"


def _categoria_financeira_do_formulario(form=None):
    """Serviço ou Produto conforme o POST (padrão: Serviço)."""
    raw = ((form.get("categoria") if form is not None else None) or "").strip()
    return "Produto" if raw == "Produto" else "Serviço"


def _id_agendamento_slot(cursor, dia, hora, barbeiro_id, barbearia_id):
    cursor.execute(
        """
        SELECT id FROM Clientes
        WHERE Dia = ? AND Hora = ? AND barbeiro_id = ? AND barbearia_id = ?
        ORDER BY id DESC LIMIT 1
        """,
        (dia, hora, barbeiro_id, barbearia_id),
    )
    row = cursor.fetchone()
    return int(row[0]) if row else None


def _listar_horarios(cursor, barbearia_id):
    cursor.execute(
        """
        SELECT hora FROM horarios
        WHERE barbearia_id = ? AND IFNULL(ativo, 1) = 1
        ORDER BY ordem, hora
        """,
        (barbearia_id,),
    )
    rows = cursor.fetchall()
    if rows:
        return [r[0] for r in rows]
    return HORARIOS


def _obter_barbearia_por_slug(cursor, slug):
    slug = (slug or "").strip()
    if not slug:
        return None
    cursor.execute(
        """
        SELECT id, nome, slug FROM barbearias
        WHERE LOWER(TRIM(slug)) = LOWER(?)
        """,
        (slug,),
    )
    return cursor.fetchone()


def _obter_barbearia_por_id(cursor, barbearia_id):
    cursor.execute(
        "SELECT id, nome, slug FROM barbearias WHERE id = ?",
        (barbearia_id,),
    )
    return cursor.fetchone()


def _barbearia_id_admin_obrigatorio():
    """Compat: ID do negócio na sessão."""
    return _barbearia_id_sessao()


def _profissional_pertence_barbearia(cursor, profissional_id, barbearia_id):
    cursor.execute(
        """
        SELECT id FROM usuarios
        WHERE id = ? AND barbearia_id = ?
          AND role IN ('barbeiro', 'profissional', 'admin')
        """,
        (profissional_id, barbearia_id),
    )
    return cursor.fetchone() is not None


def _url_segura_apos_login(next_url):
    """Evita open redirect: só paths relativos do mesmo app."""
    if not next_url:
        return None
    next_url = next_url.strip()
    if next_url.startswith("/") and not next_url.startswith("//"):
        return next_url
    return None


def _exigir_admin_ou_login():
    """Garante sessão ativa de administrador. Retorna redirect ou None."""
    return _exigir_admin()


def _identificador_e_email(identificador):
    return "@" in (identificador or "")


def _senha_confere(senha_armazenada, senha_digitada):
    if senha_armazenada and (
        senha_armazenada.startswith("pbkdf2:")
        or senha_armazenada.startswith("scrypt:")
    ):
        return check_password_hash(senha_armazenada, senha_digitada)
    return senha_armazenada == senha_digitada


def _usuarios_tem_coluna_telefone(cursor):
    return _coluna_usuarios_existe(cursor, "telefone")


def _buscar_usuario_por_identificador(cursor, identificador):
    conn = getattr(cursor, "_conn", None)
    return auth.buscar_usuario_por_identificador(identificador, conn=conn)


def _whatsapp_suporte_url(cursor):
    conn = getattr(cursor, "_conn", None)
    link = auth.whatsapp_suporte_url(conn=conn)
    return _normalizar_link_whatsapp(link) if link else None


def _url_whatsapp_texto_lembrete(nome, data, hora, servico, telefone_cliente=None):
    """Link wa.me para o profissional enviar confirmação ao cliente."""
    if not telefone_cliente:
        return None
    tel = re.sub(r"\D", "", str(telefone_cliente))
    if len(tel) < 10:
        return None
    if not tel.startswith("55"):
        tel = "55" + tel
    texto = (
        f"Olá {nome}! Seu horário está confirmado: {servico} em {data} às {hora}. "
        "Qualquer dúvida, responda aqui."
    )
    return f"https://wa.me/{tel}?text={quote(texto)}"


def _url_whatsapp_com_texto(link_base, texto):
    if not link_base:
        return None
    separador = "&" if "?" in link_base else "?"
    return f"{link_base}{separador}text={quote(texto)}"


def _link_redefinir_senha(token):
    base = (os.environ.get("APP_BASE_URL") or "").strip().rstrip("/")
    path = url_for("redefinir_senha", token=token)
    if base:
        return f"{base}{path}"
    return url_for("redefinir_senha", token=token, _external=True)


def enviar_email_recuperacao(destinatario, nome_usuario, link_redefinir):
    """
    Envia e-mail com link de redefinição via SMTP (variáveis de ambiente).
    Retorna True se enviado; False se SMTP não configurado ou falha no envio.
    """
    smtp_host = os.environ.get("SMTP_HOST", "").strip()
    smtp_port = int(os.environ.get("SMTP_PORT", "587"))
    smtp_user = os.environ.get("SMTP_USER", "").strip()
    smtp_password = os.environ.get("SMTP_PASSWORD", "").strip()
    smtp_from = os.environ.get("SMTP_FROM", smtp_user or "noreply@agendasimples.local")
    use_tls = os.environ.get("SMTP_USE_TLS", "true").lower() in ("1", "true", "yes")

    if not smtp_host:
        app.logger.warning(
            "SMTP_HOST não configurado — e-mail de recuperação não enviado para %s",
            destinatario,
        )
        return False

    assunto = "Redefinição de senha — Agenda Simples"
    corpo = f"""Olá, {nome_usuario or 'usuário'}!

Recebemos um pedido para redefinir sua senha no Agenda Simples.

Clique no link abaixo (válido por 30 minutos):
{link_redefinir}

Se você não solicitou isso, ignore este e-mail.

Atenciosamente,
Equipe Agenda Simples
"""
    msg = MIMEMultipart()
    msg["From"] = smtp_from
    msg["To"] = destinatario
    msg["Subject"] = assunto
    msg.attach(MIMEText(corpo, "plain", "utf-8"))

    try:
        with smtplib.SMTP(smtp_host, smtp_port, timeout=30) as server:
            if use_tls:
                server.starttls()
            if smtp_user and smtp_password:
                server.login(smtp_user, smtp_password)
            server.sendmail(smtp_from, [destinatario], msg.as_string())
        return True
    except Exception as exc:
        app.logger.exception("Falha ao enviar e-mail de recuperação: %s", exc)
        return False


def _inicializar_schema_aplicacao():
    """Cria tabelas/colunas conforme backend (Turso/SQLite ou SQL Server)."""
    try:
        initialize_database_schema()
    except Exception:
        app.logger.exception("Falha ao garantir schema do banco na inicialização")


init_database()
_inicializar_schema_aplicacao()

requer_plano = requer_assinatura_ativa(get_connection)
register_stripe_routes(app, get_connection)

# Chaves internas da galeria (tb_galeria.categoria) — títulos exibidos vêm do banco
CATEGORIAS_GALERIA_HOME = [
    {"slug": "corte", "modal_id": "modalCat1", "carousel_id": "carouselCat1"},
    {"slug": "corte_barba", "modal_id": "modalCat2", "carousel_id": "carouselCat2"},
    {"slug": "sobrancelha", "modal_id": "modalCat3", "carousel_id": "carouselCat3"},
    {"slug": "outros", "modal_id": "modalCat4", "carousel_id": "carouselCat4"},
]

CAPA_CATALOGO_KEYS = (
    "capa_catalogo1",
    "capa_catalogo2",
    "capa_catalogo3",
    "capa_catalogo4",
)

DEFAULT_TITULOS_CATALOGO = (
    "Catálogo 1",
    "Catálogo 2",
    "Catálogo 3",
    "Catálogo 4",
)


def _usuario_equipe_logado():
    return _role_sessao() in ROLES_EQUIPE and bool(session.get("user_id"))


def _identificador_publico_barbearia(cursor, barbearia_id=None, slug=None):
    """Slug ou ID numérico para url_for('barbearia_home', identificador=...)."""
    ident = (slug or "").strip()
    if ident:
        return ident
    if barbearia_id and cursor is not None:
        barbearia = _obter_barbearia_por_id(cursor, barbearia_id)
        if barbearia:
            slug_b = _valor_linha(barbearia, 2, "slug") or ""
            slug_b = str(slug_b).strip()
            bid = _valor_linha(barbearia, 0, "id")
            return slug_b or (str(bid) if bid is not None else None)
    return None


def _redirect_home_barbearia(cursor, barbearia_id=None, slug=None):
    """Redireciona para home.html do estabelecimento (/b/<identificador>)."""
    ident = _identificador_publico_barbearia(cursor, barbearia_id, slug)
    if ident:
        return redirect(url_for("barbearia_home", identificador=ident))
    return redirect(url_for("home"))


def _render_home_estabelecimento(cursor, barbearia_id, barbearia_slug):
    """Home pública do estabelecimento (catálogo + galeria)."""
    configs = _carregar_configs_home(cursor, barbearia_id) or {
        "nome_negocio": "AgendaSimples",
        "titulo_catalogo1": DEFAULT_TITULOS_CATALOGO[0],
        "titulo_catalogo2": DEFAULT_TITULOS_CATALOGO[1],
        "titulo_catalogo3": DEFAULT_TITULOS_CATALOGO[2],
        "titulo_catalogo4": DEFAULT_TITULOS_CATALOGO[3],
        "logotipo_url": None,
        "link_instagram": None,
        "link_facebook": None,
        "link_whatsapp": None,
    }
    cursor.execute(
        "SELECT id, categoria, caminho_foto FROM tb_galeria WHERE barbearia_id = ? ORDER BY id DESC",
        (barbearia_id,),
    )
    todas_fotos = cursor.fetchall()
    galeria = {"corte": [], "corte_barba": [], "sobrancelha": [], "outros": []}
    for f in todas_fotos:
        if f.categoria in galeria:
            galeria[f.categoria].append({"id": f.id, "foto": f.caminho_foto})
    catalogos = _montar_catalogos_home(configs, galeria)
    return render_template(
        "home.html",
        galeria=galeria,
        configs=configs,
        catalogos=catalogos,
        barbearia_slug=barbearia_slug,
    )


def _carregar_configs_home(cursor, barbearia_id):
    """Nome, logotipo, catálogos da Home e bloco direito da tela /marcar."""
    base = {
        "nome_negocio": "AgendaSimples",
        "titulo_catalogo1": DEFAULT_TITULOS_CATALOGO[0],
        "titulo_catalogo2": DEFAULT_TITULOS_CATALOGO[1],
        "titulo_catalogo3": DEFAULT_TITULOS_CATALOGO[2],
        "titulo_catalogo4": DEFAULT_TITULOS_CATALOGO[3],
        "logotipo_url": None,
        "texto_marcar_direito": None,
        "foto_fundo_direito_url": None,
        "capa_catalogo1": None,
        "capa_catalogo2": None,
        "capa_catalogo3": None,
        "capa_catalogo4": None,
        "link_instagram": None,
        "link_facebook": None,
        "link_whatsapp": None,
    }
    try:
        cursor.execute(
            """
            SELECT nome, titulo_catalogo1, titulo_catalogo2, titulo_catalogo3,
                   titulo_catalogo4, logotipo_url, texto_marcar_direito,
                   foto_fundo_direito_url, capa_catalogo1, capa_catalogo2,
                   capa_catalogo3, capa_catalogo4, link_instagram, link_facebook,
                   link_whatsapp
            FROM barbearias WHERE id = ?
            """,
            (barbearia_id,),
        )
        row = cursor.fetchone()
    except DbError:
        cursor.execute("SELECT nome FROM barbearias WHERE id = ?", (barbearia_id,))
        row = cursor.fetchone()
        if row:
            base["nome_negocio"] = row[0]
            base["texto_marcar_direito"] = row[0]
        return base

    if not row:
        return base

    titulos_db = [row[1], row[2], row[3], row[4]]
    nome = row[0] or "AgendaSimples"
    return {
        "nome_negocio": nome,
        "titulo_catalogo1": titulos_db[0] or DEFAULT_TITULOS_CATALOGO[0],
        "titulo_catalogo2": titulos_db[1] or DEFAULT_TITULOS_CATALOGO[1],
        "titulo_catalogo3": titulos_db[2] or DEFAULT_TITULOS_CATALOGO[2],
        "titulo_catalogo4": titulos_db[3] or DEFAULT_TITULOS_CATALOGO[3],
        "logotipo_url": row[5],
        "texto_marcar_direito": row[6] or nome,
        "foto_fundo_direito_url": row[7],
        "capa_catalogo1": row[8],
        "capa_catalogo2": row[9],
        "capa_catalogo3": row[10],
        "capa_catalogo4": row[11],
        "link_instagram": row[12],
        "link_facebook": row[13],
        "link_whatsapp": row[14],
    }


def _normalizar_link_whatsapp(valor):
    """Aceita URL completa ou apenas número para wa.me."""
    valor = (valor or "").strip()
    if not valor:
        return None
    if valor.startswith("http://") or valor.startswith("https://"):
        return valor
    digitos = re.sub(r"\D", "", valor)
    if not digitos:
        return None
    if not digitos.startswith("55") and len(digitos) <= 11:
        digitos = "55" + digitos
    return f"https://wa.me/{digitos}"


def _normalizar_link_url(valor):
    """Garante URL com protocolo para Instagram/Facebook."""
    valor = (valor or "").strip()
    if not valor:
        return None
    if valor.startswith("http://") or valor.startswith("https://"):
        return valor
    return f"https://{valor.lstrip('/')}"


def _processar_upload_capas_catalogo(request, cursor, barbearia_id):
    """Salva imagens de capa dos 4 catálogos da Home."""
    upload_dir = "static/uploads/catalogos"
    os.makedirs(upload_dir, exist_ok=True)
    for i, coluna in enumerate(CAPA_CATALOGO_KEYS, start=1):
        arquivo = request.files.get(coluna)
        if not arquivo or not arquivo.filename:
            continue
        ext = (
            arquivo.filename.rsplit(".", 1)[1].lower()
            if "." in arquivo.filename
            else "jpg"
        )
        filename = f"capa_catalogo{i}_{barbearia_id}.{ext}"
        arquivo.save(os.path.join(upload_dir, filename))
        try:
            cursor.execute(
                f"UPDATE barbearias SET {coluna} = ? WHERE id = ?",
                (filename, barbearia_id),
            )
        except DbError:
            pass


def _url_capa_catalogo(capa_filename, fotos_galeria):
    """URL da imagem de fundo do card: capa admin ou 1ª foto da galeria."""
    if capa_filename:
        return f"uploads/catalogos/{capa_filename}"
    if fotos_galeria:
        return f"uploads/galeria/{fotos_galeria[0]['foto']}"
    return None


def _titulo_catalogo_exibicao(titulo_raw):
    """Traduz títulos padrão (Catálogo 1…4); mantém títulos customizados do admin."""
    titulo_raw = (titulo_raw or "").strip()
    if titulo_raw in DEFAULT_TITULOS_CATALOGO:
        return _(titulo_raw)
    return titulo_raw


def _montar_catalogos_home(configs, galeria):
    catalogos = []
    titulo_keys = [
        "titulo_catalogo1",
        "titulo_catalogo2",
        "titulo_catalogo3",
        "titulo_catalogo4",
    ]
    for i, meta in enumerate(CATEGORIAS_GALERIA_HOME):
        slug = meta["slug"]
        fotos = galeria.get(slug, [])
        capa_file = configs.get(CAPA_CATALOGO_KEYS[i])
        capa_path = _url_capa_catalogo(capa_file, fotos)
        titulo_raw = configs.get(titulo_keys[i], DEFAULT_TITULOS_CATALOGO[i])
        catalogos.append({
            **meta,
            "titulo": _titulo_catalogo_exibicao(titulo_raw),
            "fotos": fotos,
            "capa_path": capa_path,
        })
    return catalogos


# Função auxiliar para buscar a foto de capa cadastrada no banco de dados
def obter_foto_capa(cursor):
    try:
        cursor.execute("SELECT valor FROM tb_configuracoes WHERE chave = 'foto_capa'")
        config_foto = cursor.fetchone()
        return config_foto[0] if config_foto else None
    except Exception as e:
        print(f"Erro ao obter foto de capa: {e}")
        return None

# Função para criar links amigáveis (SaaS)
def slugify(text):
    text = text.lower().strip()
    text = re.sub(r'[^\w\s-]', '', text)
    text = re.sub(r'[\s_-]+', '-', text)
    return text

# 3. ROTA PARA TROCAR O IDIOMA (PT | EN | ES)
@app.route("/mudar_idioma/<string:lang>")
def mudar_idioma(lang):
    if lang in app.config["BABEL_SUPPORTED_LOCALES"]:
        session["lang"] = lang
        flash(_("Idioma alterado com sucesso."), "success")
    else:
        flash(_("Idioma não suportado."), "warning")

    destino = request.args.get("next")
    if destino and destino.startswith("/") and not destino.startswith("//"):
        return redirect(destino)
    if request.referrer:
        return redirect(request.referrer)
    return redirect(url_for("home"))

# -------------------------- ROTAS DO SAAS / CADASTRO --------------------------

RAMOS_ATIVIDADE = (
    ("barbearia", _("Barbearia")),
    ("salao", _("Salão de Beleza")),
    ("estetica", _("Clínica de Estética")),
    ("outros", _("Outros")),
)
_RAMOS_VALIDOS = {c for c, _ in RAMOS_ATIVIDADE}


def _resolver_barbearia(cursor, identificador):
    """Busca estabelecimento por slug ou ID numérico."""
    conn = getattr(cursor, "_conn", None)
    if conn is not None:
        return auth.resolver_barbearia(identificador, conn=conn)
    return auth.resolver_barbearia(identificador)


def _ctx_landing_vendas(form=None):
    return {
        "trial_days": cfg.TRIAL_DAYS,
        "preco_mensal": cfg.PLANO_MENSAL_VALOR,
        "ramos": RAMOS_ATIVIDADE,
        "form": form or {},
    }


def _log_erro_cadastro_saas(exc, etapa, **contexto):
    print("=== ERRO CADASTRO SAAS ===")
    print(f"Etapa: {etapa}")
    print(f"Exceção: {type(exc).__name__}: {exc}")
    for chave, valor in contexto.items():
        print(f"  {chave}: {valor}")
    traceback.print_exc()
    print("=== FIM ERRO CADASTRO SAAS ===")
    app.logger.exception("Cadastro SaaS [%s]: %s", etapa, exc)


def _inserir_barbearia_cadastro(cursor, nome_negocio, slug, email, senha_hash, agora, ramo):
    """INSERT em barbearias com colunas opcionais conforme schema migrado."""
    base_cols = [
        "nome", "slug", "email", "senha", "plano_ativo",
        "titulo_catalogo1", "titulo_catalogo2", "titulo_catalogo3", "titulo_catalogo4",
        "texto_marcar_direito", "data_cadastro",
    ]
    base_vals = [
        nome_negocio,
        slug,
        email,
        senha_hash,
        1,
        DEFAULT_TITULOS_CATALOGO[0],
        DEFAULT_TITULOS_CATALOGO[1],
        DEFAULT_TITULOS_CATALOGO[2],
        DEFAULT_TITULOS_CATALOGO[3],
        nome_negocio,
        agora,
    ]
    if _coluna_existe(cursor, "barbearias", "ramo_atividade"):
        base_cols.append("ramo_atividade")
        base_vals.append(ramo)

    placeholders = ", ".join("?" for _ in base_vals)
    colunas_sql = ", ".join(base_cols)
    cursor.execute(
        f"INSERT INTO barbearias ({colunas_sql}) VALUES ({placeholders})",
        tuple(base_vals),
    )


def _inserir_usuario_admin_cadastro(
    cursor, nome_profissional, email, senha_hash, agora, barbearia_id
):
    """INSERT em usuarios — só inclui colunas que existem no Turso/SQLite."""
    cols = ["nome", "email", "senha", "role"]
    vals = [nome_profissional, email, senha_hash, "admin"]
    if _coluna_existe(cursor, "usuarios", "data_cadastro"):
        cols.append("data_cadastro")
        vals.append(agora)
    if _coluna_existe(cursor, "usuarios", "status_trial"):
        cols.append("status_trial")
        vals.append("trialing")
    if _coluna_existe(cursor, "usuarios", "barbearia_id"):
        cols.append("barbearia_id")
        vals.append(barbearia_id)

    placeholders = ", ".join("?" for _ in vals)
    colunas_sql = ", ".join(cols)
    cursor.execute(
        f"INSERT INTO usuarios ({colunas_sql}) VALUES ({placeholders})",
        tuple(vals),
    )


def _processar_cadastro_saas(template_name="home_vendas.html"):
    """Valida e cria conta admin + trial; retorna redirect ou template com erros."""
    conn = None
    form = {}

    try:
        nome_negocio = (request.form.get("nome_negocio") or "").strip()
        nome_profissional = (request.form.get("nome_profissional") or "").strip()
        email = (request.form.get("email") or "").strip().lower()
        ramo = (request.form.get("ramo_atividade") or "").strip().lower()
        senha = request.form.get("password") or ""
        senha_confirma = request.form.get("password_confirm") or ""

        form = {
            "nome_negocio": nome_negocio,
            "nome_profissional": nome_profissional,
            "email": email,
            "ramo_atividade": ramo,
        }

        erros = []
        if not nome_negocio:
            erros.append(_("Informe o nome do seu negócio."))
        if not nome_profissional:
            erros.append(_("Informe o seu nome."))
        if not email or "@" not in email:
            erros.append(_("Informe um e-mail válido."))
        if ramo not in _RAMOS_VALIDOS:
            erros.append(_("Selecione o ramo de atividade."))
        if len(senha) < 6:
            erros.append(_("A senha deve ter pelo menos 6 caracteres."))
        if senha != senha_confirma:
            erros.append(_("As senhas não coincidem."))

        conn, cursor = _open_db()

        if not erros and auth.email_ja_cadastrado(email, conn=conn):
            erros.append(
                _("Este e-mail já está cadastrado. Faça login ou use outro e-mail.")
            )

        slug = slugify(nome_negocio) if nome_negocio else ""
        if not erros and slug and auth.slug_ou_email_ja_usado(slug, email, conn=conn):
            erros.append(
                _(
                    "Já existe uma conta com este e-mail ou nome de negócio semelhante."
                )
            )

        if erros:
            safe_close(conn)
            for msg in erros:
                flash(msg, "error")
            return render_template(template_name, **_ctx_landing_vendas(form))

        senha_hash = generate_password_hash(senha)
        agora = datetime.utcnow().strftime("%Y-%m-%d %H:%M:%S")

        _inserir_barbearia_cadastro(
            cursor, nome_negocio, slug, email, senha_hash, agora, ramo
        )
        barbearia_id = auth.obter_id_barbearia_apos_insert(email, conn=conn)
        if not barbearia_id:
            raise RuntimeError(
                "INSERT barbearias OK, mas ID não retornado (lastrowid/Turso)."
            )

        criar_assinatura_trial(cursor, barbearia_id, cfg.TRIAL_DAYS)
        seed_servicos_horarios_padrao(cursor, barbearia_id)
        _inserir_usuario_admin_cadastro(
            cursor, nome_profissional, email, senha_hash, agora, barbearia_id
        )
        safe_commit(conn)

        user = auth.buscar_usuario_resumo_por_email(email, conn=conn)
        if not user:
            safe_close(conn)
            conn = None
            flash(
                gettext(
                    "Conta criada! Sua página pública já está no ar — faça login para gerenciá-la."
                ),
                "success",
            )
            identificador_publico = slug or str(barbearia_id)
            return redirect(
                url_for("barbearia_home", identificador=identificador_publico)
            )

        user_id = user.get("id")
        if not user_id:
            user_id = auth.obter_id_usuario_apos_insert(email, conn=conn)

        identificador_publico = slug or str(barbearia_id)
        safe_close(conn)
        conn = None

        try:
            if user_id:
                _definir_sessao_admin(
                    user_id,
                    nome_profissional or user.get("nome"),
                    barbearia_id,
                    nome_negocio,
                    slug,
                )
        except Exception as exc_sessao:
            _log_erro_cadastro_saas(
                exc_sessao,
                "sessao_pos_cadastro",
                email=email,
                barbearia_id=barbearia_id,
            )

        flash(
            gettext(
                "Parabéns! Sua conta teste de %(days)s dias foi criada. "
                "Esta é a página pública do seu negócio — já está no ar para seus clientes!"
            )
            % {"days": int(cfg.TRIAL_DAYS)},
            "success",
        )
        return redirect(
            url_for("barbearia_home", identificador=identificador_publico)
        )

    except Exception as exc:
        safe_rollback(conn)
        _log_erro_cadastro_saas(
            exc,
            "processamento_cadastro",
            email=form.get("email"),
            negocio=form.get("nome_negocio"),
        )
        flash(
            _("Não foi possível criar sua conta. Tente novamente em instantes."),
            "error",
        )
        return render_template(template_name, **_ctx_landing_vendas(form))

    finally:
        safe_close(conn)


@app.route("/registrar", methods=["GET", "POST"])
@app.route("/cadastro", methods=["GET", "POST"])
def registrar():
    """Cadastro SaaS — GET redireciona à landing; POST processa o formulário."""
    if request.method == "GET":
        if _usuario_autenticado() and _role_sessao() in ROLES_EQUIPE:
            return redirect(url_for("admin_agenda"))
        return redirect(url_for("home") + "#cadastro")
    return _processar_cadastro_saas("home_vendas.html")


@app.route("/cadastro_barbearia", methods=["GET", "POST"])
def cadastro_barbearia():
    """Legado — redireciona para a landing de vendas."""
    return redirect(url_for("home") + "#cadastro")


# -------------------------- LANDING DE VENDAS (/) --------------------------
@app.route("/", methods=["GET", "POST"])
def home():
    """Landing Page SaaS — conversão para plano mensal."""
    if request.method == "POST":
        return _processar_cadastro_saas("home_vendas.html")
    if _usuario_autenticado() and _role_sessao() in ROLES_EQUIPE:
        return redirect(url_for("admin_agenda"))
    return render_template("home_vendas.html", **_ctx_landing_vendas())


@app.route("/b/<identificador>")
def barbearia_home(identificador):
    """Home pública do estabelecimento (slug ou ID numérico)."""
    with db_adapter.connection_scope() as conn:
        db_adapter.ensure_app_schema(conn)
        cursor = db_adapter.cursor(conn)
        barbearia = auth.resolver_barbearia(identificador, conn=conn)
        if not barbearia:
            flash(_("Estabelecimento não encontrado."), "warning")
            return redirect(url_for("home"))
        slug_exib = barbearia["slug"] or str(barbearia["id"])
        return _render_home_estabelecimento(cursor, barbearia["id"], slug_exib)


@app.route("/assinatura/bloqueio")
def bloqueio_assinatura():
    """Tela de bloqueio quando trial expirou e plano não está ativo."""
    bloqueio = _exigir_login_barbearia()
    if bloqueio:
        return bloqueio
    barbearia_id = session.get("barbearia_id")
    if not barbearia_id:
        return redirect(url_for("login"))

    conn, cursor = _open_db()
    from subscriptions import admin_tem_acesso_painel

    if admin_tem_acesso_painel(cursor, barbearia_id):
        safe_close(conn)
        return redirect(url_for("admin_agenda"))
    safe_close(conn)
    return render_template(
        "bloqueio_assinatura.html",
        preco=cfg.PLANO_MENSAL_VALOR,
    )


# -------------------------- LOGIN / LOGOUT --------------------------

@app.route("/login", methods=["GET", "POST"])
def login():
    if request.method == "POST":
        email = request.form.get("email")
        senha = request.form.get("password")
        conn, cursor = _open_db()
        try:
            usuario = auth.buscar_usuario_por_email(email, conn=conn)

            senha_armazenada = (usuario or {}).get("senha")
            if usuario and _senha_confere(senha_armazenada, senha):
                user_id = int(usuario["id"])
                bid_db = auth.barbearia_id_do_usuario(
                    user_id, conn=conn, tentar_corrigir=True
                )
                if bid_db is None:
                    safe_rollback(conn)
                    app.logger.warning(
                        "Login bloqueado: user_id=%s sem barbearia_id no banco",
                        user_id,
                    )
                    return _logout_por_falta_vinculo()

                try:
                    usuario["barbearia_id"] = bid_db
                    _iniciar_sessao_usuario(cursor, usuario, email)
                    session["barbearia_id"] = int(bid_db)
                    session["usuario_id"] = user_id
                    session["user_id"] = user_id
                    session.modified = True

                    bid_final = auth.barbearia_id_do_usuario(
                        user_id, conn=conn, tentar_corrigir=False
                    )
                    if bid_final is None or not session.get("barbearia_id"):
                        safe_rollback(conn)
                        return _logout_por_falta_vinculo()

                    app.logger.info(
                        "Login OK user_id=%s barbearia_id=%s",
                        user_id,
                        bid_final,
                    )
                    safe_commit(conn)
                except ValueError as exc:
                    safe_rollback(conn)
                    app.logger.warning("Login sem barbearia vinculada: %s", exc)
                    return _logout_por_falta_vinculo()

                destino = _url_segura_apos_login(
                    request.form.get("next") or request.args.get("next")
                )
                if destino:
                    return redirect(destino)
                return redirect(url_for("admin_agenda"))

            safe_rollback(conn)
        except Exception:
            safe_rollback(conn)
            app.logger.exception("Erro no login")
            flash(_("Não foi possível concluir o login. Tente novamente."), "error")
            return redirect(url_for("login", next=request.form.get("next")))
        finally:
            safe_close(conn)

        flash("Usuário ou senha incorretos!", "error")
        return redirect(url_for("login", next=request.form.get("next")))
    proxima = request.args.get("next")
    if proxima and _url_segura_apos_login(proxima):
        flash(_("Faça login para acessar esta página."), "warning")
    return render_template("login.html", next_url=proxima)


@app.route("/esqueci_senha", methods=["GET", "POST"])
def esqueci_senha():
    whatsapp_url = None
    telefone_exibicao = None
    modo_whatsapp = False

    if request.method == "POST":
        identificador = (request.form.get("identificador") or "").strip()
        usuario = auth.buscar_usuario_por_identificador(identificador)

        if not usuario:
            flash(
                "Não encontramos cadastro com esse e-mail ou telefone. Verifique os dados.",
                "error",
            )
            return render_template("esqueci_senha.html")

        if _identificador_e_email(identificador):
            token = _gerar_token_recuperacao(usuario["id"])
            link = _link_redefinir_senha(token)
            enviado = enviar_email_recuperacao(
                usuario["email"], usuario["nome"], link
            )
            if enviado:
                flash(
                    "Enviamos um link de recuperação para o seu e-mail. "
                    "O link expira em 30 minutos.",
                    "success",
                )
            else:
                flash(
                    "Conta encontrada, mas não foi possível enviar o e-mail. "
                    "Configure SMTP no servidor ou use seu telefone cadastrado. "
                    "Em ambiente de desenvolvimento, verifique os logs do servidor.",
                    "warning",
                )
            return render_template("esqueci_senha.html")

        telefone_limpo = _limpar_telefone(identificador)
        telefone_exibicao = identificador
        texto = (
            f"Olá, esqueci minha senha e gostaria de redefinir. "
            f"Meu telefone é {telefone_exibicao}."
        )
        wa_base = auth.whatsapp_suporte_url()
        if wa_base:
            wa_base = _normalizar_link_whatsapp(wa_base)
        whatsapp_url = _url_whatsapp_com_texto(wa_base, texto)
        modo_whatsapp = True
        if not whatsapp_url:
            flash(
                "Encontramos seu cadastro, mas o WhatsApp de suporte ainda não está "
                "configurado. Entre em contato com o administrador da barbearia.",
                "warning",
            )
        return render_template(
            "esqueci_senha.html",
            modo_whatsapp=modo_whatsapp,
            whatsapp_url=whatsapp_url,
            telefone_exibicao=telefone_exibicao,
        )

    return render_template(
        "esqueci_senha.html",
        modo_whatsapp=modo_whatsapp,
        whatsapp_url=whatsapp_url,
        telefone_exibicao=telefone_exibicao,
    )


@app.route("/redefinir_senha/<token>", methods=["GET", "POST"])
def redefinir_senha(token):
    user_id = _validar_token_recuperacao(token)
    if not user_id:
        flash(
            "Link inválido ou expirado. Solicite uma nova recuperação de senha.",
            "error",
        )
        return redirect(url_for("esqueci_senha"))

    if request.method == "POST":
        nova = request.form.get("nova_senha", "")
        confirma = request.form.get("confirma_senha", "")
        if len(nova) < 6:
            flash("A senha deve ter pelo menos 6 caracteres.", "error")
            return render_template("redefinir_senha.html", token=token)
        if nova != confirma:
            flash("As senhas não coincidem.", "error")
            return render_template("redefinir_senha.html", token=token)

        senha_hash = generate_password_hash(nova)
        auth.atualizar_senha_usuario(user_id, senha_hash)
        flash("Senha alterada com sucesso! Faça login com a nova senha.", "success")
        return redirect(url_for("login"))

    return render_template("redefinir_senha.html", token=token)


@app.route("/logout")
def logout():
    """Encerra sessão e volta à página pública do negócio (não à landing SaaS)."""
    identificador = (session.get("barbearia_slug") or "").strip()
    if not identificador and session.get("barbearia_id"):
        identificador = str(session.get("barbearia_id"))

    session.clear()

    if identificador:
        return redirect(url_for("barbearia_home", identificador=identificador))
    return redirect(url_for("home"))

# -------------------------- AGENDA BARBEIRO (legado → painel unificado) --------------------------
@app.route("/agenda")
def agenda():
    bloqueio = _exigir_login_barbearia()
    if bloqueio:
        return bloqueio
    return redirect(url_for("admin_agenda"))


@app.route("/acesso-negado")
def acesso_negado():
    identificador = (session.get("barbearia_slug") or "").strip()
    if not identificador and session.get("barbearia_id"):
        identificador = str(session.get("barbearia_id"))
    return render_template(
        "acesso_negado.html",
        identificador=identificador,
        role=_role_sessao(),
    )


# -------------------------- AGENDA DO NEGÓCIO (admin + profissionais) --------------------------
@app.route("/admin")
@app.route("/admin_agenda")
@requer_plano
def admin_agenda():
    bloqueio = _exigir_login_barbearia()
    if bloqueio:
        return bloqueio

    try:
        with db_adapter.connection_scope() as conn:
            db_adapter.ensure_app_schema(conn)
            cursor = db_adapter.cursor(conn)
            bloqueio_vinculo = _sincronizar_barbearia_sessao(cursor)
            if bloqueio_vinculo:
                return bloqueio_vinculo

            barbearia_id = _barbearia_id_sessao()
            user_id = session.get("user_id") or session.get("usuario_id")
            eh_admin = _eh_admin()
            filtrar_meus = _eh_profissional_equipe()

            foto_capa = obter_foto_capa(cursor)
            horarios_negocio = agenda.listar_horarios(barbearia_id, conn=conn)
            registros = agenda.listar_registros_agenda(
                barbearia_id, user_id, filtrar_meus, conn=conn
            )

            hoje = datetime.today()
            agenda_data = {
                (hoje + timedelta(days=i)).strftime("%Y-%m-%d"): {
                    h: None for h in horarios_negocio
                }
                for i in range(28)
            }

            for r in registros:
                dia_val = r.get("Dia")
                hora_val = r.get("Hora")
                d_str = (
                    dia_val.strftime("%Y-%m-%d")
                    if isinstance(dia_val, datetime)
                    else str(dia_val or "")[:10]
                )
                h_str = str(hora_val or "")[:5]
                if d_str in agenda_data and h_str in agenda_data[d_str]:
                    if agenda_data[d_str][h_str] is None:
                        agenda_data[d_str][h_str] = []
                    agenda_data[d_str][h_str].append({
                        "nome": r.get("Nome"),
                        "servico": r.get("Servico"),
                        "whatsapp": r.get("Whatsapp") or "",
                        "barbeiro_nome": r.get("barbeiro_nome"),
                        "barbeiro_id": r.get("barbeiro_id"),
                    })

            barbeiros = fin.listar_profissionais(barbearia_id, conn=conn)
            barbearia = auth.buscar_barbearia_por_id(barbearia_id, conn=conn)
            barbearia_slug = (
                barbearia["slug"] if barbearia else session.get("barbearia_slug", "")
            )
    except Exception as exc:
        app.logger.exception("Erro ao carregar admin_agenda: %s", exc)
        flash(_("Não foi possível carregar a agenda. Tente novamente."), "error")
        return redirect(url_for("acesso_negado"))

    return render_template(
        "admin_agenda.html",
        agenda=agenda_data,
        horarios=horarios_negocio,
        datetime=datetime,
        dias_pt=DIAS_PT,
        barbeiros=barbeiros,
        foto_capa=foto_capa,
        barbearia_slug=barbearia_slug,
        eh_admin=eh_admin,
        filtrar_meus=filtrar_meus,
    )


@app.route("/admin/profissionais/novo", methods=["GET", "POST"])
@app.route("/admin_cadastrar_profissional", methods=["GET", "POST"])
@app.route("/admin/profissional/novo", methods=["GET", "POST"])
@requer_plano
def admin_cadastrar_profissional():
    bloqueio = _exigir_admin_ou_login()
    if bloqueio:
        return bloqueio

    if request.method == "POST":
        barbearia_id_sessao = _barbearia_id_admin_sessao_obrigatorio()
        if not barbearia_id_sessao:
            flash(
                _(
                    "Não foi possível cadastrar o profissional: negócio não identificado "
                    "na sessão. Faça login novamente como administrador."
                ),
                "error",
            )
            return redirect(url_for("login"))

        nome = (request.form.get("nome") or "").strip()
        email = (request.form.get("email") or "").strip()
        celular_raw = (request.form.get("celular") or "").strip()
        especialidade = (request.form.get("especialidade") or "").strip()
        senha = request.form.get("senha") or ""
        foto = request.files.get("foto_perfil")

        erros = []
        if not nome:
            erros.append("Informe o nome completo.")
        if not email or "@" not in email:
            erros.append("Informe um e-mail válido.")
        if len(senha) < 6:
            erros.append("A senha inicial deve ter pelo menos 6 caracteres.")

        celular = _limpar_telefone(celular_raw) if celular_raw else None
        if celular_raw and celular and len(celular) < 10:
            erros.append("Celular inválido.")

        foto_nome = None
        if foto and foto.filename:
            foto_nome = _salvar_upload_foto_perfil(foto)
            if not foto_nome:
                erros.append(
                    "Foto inválida. Use JPG, PNG, WEBP ou GIF (máx. recomendado 5 MB)."
                )

        barbearia_id = barbearia_id_sessao

        app.logger.info(
            "Cadastro profissional: admin user_id=%s barbearia_id sessão=%s",
            session.get("user_id"),
            barbearia_id,
        )

        conn, cursor = _open_db()

        if not _coluna_existe(cursor, "usuarios", "barbearia_id"):
            safe_close(conn)
            flash(
                _(
                    "Banco desatualizado. Execute: python aplicar_migracao_barbearia_id_usuarios.py"
                ),
                "error",
            )
            return render_template("admin_cadastrar_profissional.html", form={})

        if email and not erros and auth.email_ja_cadastrado(email, conn=conn):
            erros.append("Este e-mail já está cadastrado.")

        if erros:
            safe_close(conn)
            for msg in erros:
                flash(msg, "error")
            return render_template(
                "admin_cadastrar_profissional.html",
                form={
                    "nome": nome,
                    "email": email,
                    "celular": celular_raw,
                    "especialidade": especialidade,
                },
            )

        senha_hash = generate_password_hash(senha)

        try:
            _inserir_profissional_usuario(
                cursor,
                nome,
                email,
                senha_hash,
                barbearia_id,
                telefone=celular or None,
                especialidade=especialidade or None,
                foto_perfil=foto_nome,
            )
        except ValueError as exc:
            safe_close(conn)
            flash(str(exc), "error")
            return render_template(
                "admin_cadastrar_profissional.html",
                form={
                    "nome": nome,
                    "email": email,
                    "celular": celular_raw,
                    "especialidade": especialidade,
                },
            )

        ver = db_adapter.execute_query(
            "SELECT barbearia_id FROM usuarios WHERE LOWER(TRIM(email)) = LOWER(?)",
            (email,),
            fetch="one",
            conn=conn,
        )
        bid_inserido = db_adapter.row_get(ver, "barbearia_id", index=0) if ver else None
        if bid_inserido is None or int(bid_inserido) != int(barbearia_id):
            safe_rollback(conn)
            safe_close(conn)
            flash(
                _(
                    "Falha ao gravar vínculo com o negócio. "
                    "O profissional não foi cadastrado. Tente novamente."
                ),
                "error",
            )
            return render_template("admin_cadastrar_profissional.html", form={})

        app.logger.info(
            "Profissional criado nome=%s email=%s barbearia_id=%s",
            nome,
            email,
            barbearia_id,
        )
        safe_commit(conn)
        safe_close(conn)
        flash(f"Profissional {nome} cadastrado com sucesso!", "success")
        return redirect(url_for("admin_agenda"))

    return render_template("admin_cadastrar_profissional.html", form={})


@app.route("/admin/agenda/concluir", methods=["POST"])
@requer_plano
def admin_agenda_concluir():
    bloqueio = _exigir_login_barbearia()
    if bloqueio:
        return bloqueio

    data = request.form.get("data")
    hora = request.form.get("hora")
    barbeiro_id = request.form.get("barbeiro_id")
    valor_raw = (request.form.get("valor") or "0").strip().replace(",", ".")
    tipo_feito = request.form.get("descricao_tipo", "Serviço")
    detalhe = request.form.get("descricao_detalhe", "")
    servico_agendado = request.form.get("servico_agendado", "")
    nome_cliente = request.form.get("nome_cliente", "")

    if not data or not hora or not barbeiro_id:
        flash("Dados do agendamento incompletos.", "danger")
        return redirect(url_for("admin_agenda"))

    try:
        valor = float(valor_raw) if valor_raw else 0.0
    except ValueError:
        flash("Valor inválido. Use apenas números (ex: 50 ou 50.00).", "danger")
        return redirect(url_for("admin_agenda"))

    if valor < 0:
        flash("O valor não pode ser negativo.", "warning")
        return redirect(url_for("admin_agenda"))

    try:
        barbeiro_id_int = int(barbeiro_id)
    except (TypeError, ValueError):
        flash("Profissional inválido.", "danger")
        return redirect(url_for("admin_agenda"))

    if _eh_profissional_equipe() and barbeiro_id_int != int(session["user_id"]):
        flash(_("Você só pode concluir seus próprios agendamentos."), "error")
        return redirect(url_for("acesso_negado"))

    barbearia_id = _barbearia_id_sessao()

    try:
        with db_adapter.connection_scope() as conn:
            db_adapter.ensure_app_schema(conn)
            cursor = db_adapter.cursor(conn)
            agendamento = agenda.buscar_agendamento_slot(
                data, hora, int(barbeiro_id), barbearia_id, conn=conn
            )
            if not agendamento:
                flash("Agendamento não encontrado.", "danger")
                return redirect(url_for("admin_agenda"))
            if agendamento.get("status") == "Concluído":
                flash("Este agendamento já foi concluído.", "warning")
                return redirect(url_for("admin_agenda"))

            agenda.marcar_agendamento_concluido(
                data, hora, int(barbeiro_id), barbearia_id, conn=conn
            )
            agendamento_id = agenda.id_agendamento_slot(
                data, hora, int(barbeiro_id), barbearia_id, conn=conn
            )
            descricao_fin = _montar_descricao_atendimento(
                tipo_feito, detalhe, servico_agendado, nome_cliente
            )
            cat_fin = _categoria_de_tipo_atendimento(tipo_feito)
            try:
                if valor > 0:
                    _registrar_receita_agendamento(
                        cursor,
                        descricao_fin,
                        valor,
                        int(barbeiro_id),
                        barbearia_id,
                        agendamento_id=agendamento_id,
                        substituir_existente=True,
                        categoria=cat_fin,
                    )
                    flash(
                        f"Atendimento concluído. Receita de R$ {valor:.2f} registrada no financeiro.",
                        "success",
                    )
                else:
                    if agendamento_id:
                        _registrar_receita_agendamento(
                            cursor,
                            descricao_fin,
                            0.0,
                            int(barbeiro_id),
                            barbearia_id,
                            agendamento_id=agendamento_id,
                            substituir_existente=True,
                            categoria=cat_fin,
                        )
                    flash(
                        "Atendimento concluído sem lançamento financeiro (valor R$ 0,00).",
                        "success",
                    )
            except Exception:
                app.logger.exception(
                    "Falha no financeiro ao concluir %s %s barbeiro=%s",
                    data,
                    hora,
                    barbeiro_id,
                )
                flash(
                    _(
                        "Atendimento marcado como concluído, mas o lançamento "
                        "financeiro falhou. Registre manualmente no financeiro."
                    ),
                    "warning",
                )
    except DbError as e:
        flash(f"Erro ao concluir atendimento: {e}", "danger")
    except Exception:
        app.logger.exception("Erro ao concluir atendimento")
        flash(_("Não foi possível concluir o atendimento."), "danger")

    return redirect(url_for("admin_agenda"))


# -------------------------- AGENDAR --------------------------
def _normalizar_id_inserido(valor):
    """Converte ID retornado pelo SQL Server (int, Decimal, None)."""
    if valor is None:
        return None
    try:
        return int(valor)
    except (TypeError, ValueError):
        try:
            return int(float(valor))
        except (TypeError, ValueError):
            return None


def _inserir_agendamento_retornar_id(
    cursor,
    nome,
    data,
    hora,
    servico,
    whatsapp,
    barbeiro_id,
    barbearia_id,
    valor=0.0,
):
    """Insere agendamento na tabela Clientes e devolve o id gerado."""
    valor_gravado = float(valor if valor is not None else 0)
    cursor.execute(
        """
        INSERT INTO Clientes (
            Nome, Dia, Hora, Servico, Whatsapp,
            barbeiro_id, barbearia_id, status, valor
        )
        VALUES (?, ?, ?, ?, ?, ?, ?, 'Agendado', ?)
        """,
        (
            nome,
            data,
            hora,
            servico,
            whatsapp,
            barbeiro_id,
            barbearia_id,
            valor_gravado,
        ),
    )
    return _normalizar_id_inserido(cursor.lastrowid)


def _buscar_agendamento_por_id(cursor, agendamento_id):
    """Retorna dict com dados do agendamento ou None."""
    cursor.execute(
        """
        SELECT c.id, c.Nome, c.Dia, c.Hora, c.Servico, c.Whatsapp,
               c.barbeiro_id, u.nome AS profissional_nome,
               c.barbearia_id, b.slug AS barbearia_slug
        FROM Clientes c
        LEFT JOIN usuarios u ON c.barbeiro_id = u.id
        LEFT JOIN barbearias b ON c.barbearia_id = b.id
        WHERE c.id = ?
        """,
        (agendamento_id,),
    )
    row = cursor.fetchone()
    if not row:
        return None

    dia = row[2]
    if hasattr(dia, "strftime"):
        data_fmt = dia.strftime("%Y-%m-%d")
        data_exib = dia.strftime("%d/%m/%Y")
    else:
        data_fmt = str(dia)[:10]
        data_exib = data_fmt

    hora = str(row[3])[:5] if row[3] else ""

    return {
        "id": row[0],
        "nome": row[1],
        "data": data_fmt,
        "data_exib": data_exib,
        "hora": hora,
        "servico": row[4],
        "whatsapp": row[5] or "",
        "barbeiro_id": row[6],
        "profissional": row[7] or "—",
        "barbearia_id": row[8],
        "barbearia_slug": (row[9] or "").strip() if len(row) > 9 else "",
    }


@app.route("/agendar", methods=["POST"])
def agendar():
    nome = (request.form.get("nome") or "").strip()
    data = (request.form.get("data") or "").strip()
    hora = (request.form.get("hora") or "").strip()
    servico = (request.form.get("servico") or "").strip()
    whatsapp = request.form.get("whatsapp", "")
    barbeiro_id_raw = (request.form.get("barbeiro_id") or "").strip()
    barbearia_id = request.form.get("barbearia_id", type=int)
    slug_volta = (request.form.get("barbearia_slug") or "").strip()
    if not barbearia_id and session.get("barbearia_id"):
        barbearia_id = int(session["barbearia_id"])
    if not slug_volta and session.get("barbearia_slug"):
        slug_volta = str(session["barbearia_slug"]).strip()

    campos_obrigatorios = {
        "nome": nome,
        "data": data,
        "hora": hora,
        "servico": servico,
        "barbeiro_id": barbeiro_id_raw,
    }
    if not all(campos_obrigatorios.values()):
        flash(_("Preencha todos os campos obrigatórios do agendamento."), "error")
        return redirect(url_for("home"))

    try:
        barbeiro_id = int(barbeiro_id_raw)
    except (TypeError, ValueError):
        flash(_("Profissional inválido."), "error")
        return redirect(url_for("home"))

    novo_id = None
    ident_home = slug_volta or None
    equipe = False
    try:
        with db_adapter.connection_scope() as conn:
            db_adapter.ensure_app_schema(conn)
            cursor = db_adapter.cursor(conn)
            equipe = _usuario_equipe_logado()

            if not barbearia_id and slug_volta:
                barbearia = auth.buscar_barbearia_por_slug(slug_volta, conn=conn)
                if barbearia:
                    barbearia_id = barbearia["id"]

            if not barbearia_id:
                flash(_("Estabelecimento inválido."), "error")
                return redirect(url_for("home"))

            if not agenda.profissional_pertence_barbearia(
                barbeiro_id, barbearia_id, conn=conn
            ):
                flash(_("Profissional inválido para este estabelecimento."), "error")
                if equipe:
                    return redirect(url_for("admin_agenda"))
                return _redirect_home_barbearia(cursor, barbearia_id, slug_volta)

            if agenda.horario_ocupado(
                data, hora, barbeiro_id, barbearia_id, conn=conn
            ):
                flash(
                    "Este horário já está ocupado com este profissional. Por favor, escolha outra opção!",
                    "warning",
                )
                if equipe:
                    return redirect(url_for("admin_agenda"))
                return _redirect_home_barbearia(cursor, barbearia_id, slug_volta)

            valor_form = _parse_valor_monetario(request.form.get("valor"))
            valor_clientes = float(valor_form) if valor_form is not None else 0.0
            valor_financeiro = (
                float(valor_form)
                if valor_form is not None
                else agenda.obter_preco_servico(barbearia_id, servico, conn=conn)
            )
            categoria_fin = _categoria_financeira_do_formulario(request.form)
            descricao_fin = _descricao_financeiro_agendamento(servico, nome)

            novo_id = agenda.inserir_agendamento(
                nome,
                data,
                hora,
                servico,
                whatsapp,
                barbeiro_id,
                barbearia_id,
                valor=valor_clientes,
                conn=conn,
            )
            ident_home = _identificador_publico_barbearia(
                cursor, barbearia_id, slug_volta
            )

            if novo_id:
                try:
                    _registrar_receita_agendamento(
                        cursor,
                        descricao_fin,
                        valor_financeiro,
                        barbeiro_id,
                        barbearia_id,
                        agendamento_id=novo_id,
                        substituir_existente=True,
                        categoria=categoria_fin,
                        nome_item=servico,
                    )
                except Exception:
                    app.logger.exception(
                        "Falha ao registrar financeiro (agendamento id=%s já salvo)",
                        novo_id,
                    )
    except Exception as exc:
        app.logger.exception("Erro ao salvar agendamento: %s", exc)
        flash(
            _("Não foi possível salvar o agendamento. Detalhe: %(erro)s", erro=str(exc)),
            "danger",
        )
        if equipe:
            return redirect(url_for("admin_agenda"))
        if ident_home:
            return redirect(url_for("barbearia_home", identificador=ident_home))
        return redirect(url_for("home"))

    if not novo_id:
        flash(
            "Agendamento salvo, mas não foi possível abrir a tela de confirmação. "
            "Verifique se a coluna id existe em Clientes (migracao_clientes_id).",
            "warning",
        )
        if equipe:
            return redirect(url_for("admin_agenda"))
        if ident_home:
            return redirect(url_for("barbearia_home", identificador=ident_home))
        return redirect(url_for("home"))

    if equipe:
        return redirect(
            url_for(
                "sucesso_agendamento",
                agendamento_id=novo_id,
                slug=ident_home,
            )
        )

    flash(
        _(
            "Agendamento confirmado! %(nome)s — %(data)s às %(hora)s (%(servico)s).",
            nome=nome,
            data=data,
            hora=hora,
            servico=servico,
        ),
        "success",
    )
    if ident_home:
        return redirect(url_for("barbearia_home", identificador=ident_home))
    return redirect(url_for("home"))


@app.route("/sucesso/<int:agendamento_id>")
def sucesso_agendamento(agendamento_id):
    slug = request.args.get("slug")
    if not agendamento_id or agendamento_id < 1:
        flash("Link de confirmação inválido.", "warning")
        if slug:
            return redirect(url_for("barbearia_home", identificador=slug))
        return redirect(url_for("home"))

    try:
        ag = agenda.buscar_agendamento_por_id(agendamento_id)
    except DbError:
        ag = None

    if not ag:
        flash("Agendamento não encontrado ou já removido.", "danger")
        if slug:
            return redirect(url_for("barbearia_home", identificador=slug))
        return redirect(url_for("home"))

    ident_home = slug or ag.get("barbearia_slug") or ""
    if not ident_home and ag.get("barbearia_id"):
        b = auth.buscar_barbearia_por_id(int(ag["barbearia_id"]))
        if b:
            ident_home = (b.get("slug") or str(b["id"])).strip()

    if not _usuario_equipe_logado():
        if ident_home:
            flash(_("Seu horário foi marcado com sucesso!"), "success")
            return redirect(url_for("barbearia_home", identificador=ident_home))
        return redirect(url_for("home"))

    url_home = (
        url_for("barbearia_home", identificador=ident_home)
        if ident_home
        else url_for("home")
    )

    whatsapp_lembrete = _url_whatsapp_texto_lembrete(
        ag.get("nome"),
        ag.get("data_exib") or ag.get("data"),
        ag.get("hora"),
        ag.get("servico"),
        ag.get("whatsapp"),
    )

    return render_template(
        "sucesso_agendamento.html",
        ag=ag,
        voltar_para_agenda=True,
        barbearia_slug=ident_home,
        url_home=url_home,
        whatsapp_lembrete=whatsapp_lembrete,
    )

# -------------------------- WHATSAPP / EDITAR / EXCLUIR --------------------------
@app.route("/editar/<string:data>/<string:hora>", methods=["GET", "POST"])
def editar(data, hora):
    bloqueio = _exigir_login_barbearia()
    if bloqueio:
        return bloqueio
    barbearia_id = _barbearia_id_sessao()
    barbeiro_id = request.args.get("barbeiro_id") or request.form.get("barbeiro_id")
    if _eh_profissional_equipe() and str(barbeiro_id) != str(session.get("user_id")):
        flash(_("Você só pode editar seus próprios agendamentos."), "error")
        return redirect(url_for("acesso_negado"))
    if request.method == "POST":
        agenda.atualizar_cliente_slot(
            request.form["nome"],
            request.form["servico"],
            request.form.get("whatsapp", ""),
            data,
            hora,
            int(barbeiro_id),
            barbearia_id,
        )
        return redirect(url_for("admin_agenda"))
    cliente_row = agenda.buscar_cliente_slot(
        data, hora, int(barbeiro_id), barbearia_id
    )
    cliente = (
        (cliente_row["Nome"], cliente_row["Servico"], cliente_row["Whatsapp"])
        if cliente_row
        else None
    )
    return render_template("editar.html", cliente=cliente, data=data, hora=hora, barbeiro_id=barbeiro_id)

@app.route("/excluir/<string:data>/<string:hora>")
def excluir(data, hora):
    bloqueio = _exigir_login_barbearia()
    if bloqueio:
        return bloqueio
    barbearia_id = _barbearia_id_sessao()
    barbeiro_id = request.args.get("barbeiro_id")
    if _eh_profissional_equipe() and str(barbeiro_id) != str(session.get("user_id")):
        flash(_("Você só pode excluir seus próprios agendamentos."), "error")
        return redirect(url_for("acesso_negado"))
    agenda.excluir_cliente_slot(data, hora, int(barbeiro_id), barbearia_id)
    return redirect(url_for("admin_agenda"))

@app.route("/whatsapp/<string:data>/<string:hora>")
def enviar_whatsapp(data, hora):
    bloqueio = _exigir_login_barbearia()
    if bloqueio:
        return bloqueio
    barbearia_id = _barbearia_id_sessao()
    barbeiro_id = request.args.get("barbeiro_id")
    if _eh_profissional_equipe() and str(barbeiro_id) != str(session.get("user_id")):
        return redirect(url_for("acesso_negado"))
    cliente = agenda.buscar_whatsapp_cliente_slot(
        data, hora, int(barbeiro_id), barbearia_id
    )
    if not cliente or not cliente[1]:
        return "⚠️ WhatsApp não cadastrado!"
    numero = "+55" + cliente[1].strip() if not cliente[1].startswith("+") else cliente[1].strip()
    try:
        import pywhatkit

        pywhatkit.sendwhatmsg_instantly(
            numero,
            f"Lembrete: {cliente[0]}, seu horário é {data} às {hora}.",
            wait_time=15,
        )
        return "✅ Enviado!"
    except Exception:
        return "❌ Erro ao enviar."

# -------------------------- MARCAR AGENDAMENTO CLIENTE --------------------------
@app.route("/marcar")
def marcar():
    """Legado: redireciona para a URL pública do estabelecimento."""
    slug_sessao = session.get("barbearia_slug")
    if slug_sessao:
        return redirect(url_for("marcar_barbearia", identificador=slug_sessao, **request.args))
    slug = auth.primeiro_slug_barbearia()
    if slug:
        return redirect(url_for("marcar_barbearia", identificador=slug, **request.args))
    flash(_("Nenhum estabelecimento cadastrado ainda."), "warning")
    return redirect(url_for("home"))


@app.route("/b/<identificador>/marcar")
def marcar_barbearia(identificador):
    conn, cursor = _open_db()
    barbearia = auth.resolver_barbearia(identificador, conn=conn)
    if not barbearia:
        safe_close(conn)
        flash(_("Estabelecimento não encontrado."), "warning")
        return redirect(url_for("home"))

    barbearia_id = barbearia["id"]
    slug_exib = barbearia["slug"] or str(barbearia["id"])
    configs = _carregar_configs_home(cursor, barbearia_id)
    barbeiros = fin.listar_profissionais(barbearia_id, conn=conn)
    horarios = agenda.listar_horarios(barbearia_id, conn=conn)
    servicos = _listar_servicos(cursor, barbearia_id)
    servicos_detalhados = _listar_servicos_detalhados(cursor, barbearia_id)
    safe_close(conn)

    profissional_sugerido = None
    profissional_param = request.args.get("profissional")
    if profissional_param:
        try:
            pid = int(profissional_param)
            ids_validos = {b[0] for b in barbeiros}
            if pid in ids_validos:
                profissional_sugerido = pid
        except (TypeError, ValueError):
            pass

    return render_template(
        "marcar_agendamento.html",
        barbeiros=barbeiros,
        horarios=horarios,
        servicos=servicos,
        servicos_detalhados=servicos_detalhados,
        profissional_sugerido=profissional_sugerido,
        configs=configs,
        barbearia_id=barbearia_id,
        barbearia_slug=slug_exib,
    )

# -------------------------- EXPORTAR PDF / EXCEL --------------------------
@app.route("/exportar_excel")
def exportar_excel():
    bloqueio = _exigir_admin()
    if bloqueio:
        return bloqueio
    barbearia_id = _barbearia_id_admin_obrigatorio()
    if not barbearia_id:
        return redirect(url_for("login"))

    dados = agenda.listar_agendamentos_export(barbearia_id)

    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.append(["Nome", "Data", "Hora", "Serviço"])
    for d in dados:
        ws.append([d[0], str(d[1]), d[2], d[3]])

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return send_file(output, download_name="agenda.xlsx", as_attachment=True)

@app.route("/pdf_hoje")
def pdf_hoje():
    return redirect(url_for("pdf_diario", data=datetime.today().strftime("%Y-%m-%d")))

@app.route("/pdf_diario/<string:data>")
def pdf_diario(data):
    bloqueio = _exigir_admin()
    if bloqueio:
        return bloqueio
    barbearia_id = _barbearia_id_sessao()
    clientes = agenda.listar_agendamentos_dia_pdf(data, barbearia_id)

    output = io.BytesIO()
    c = canvas.Canvas(output, pagesize=A4)
    c.drawString(50, 800, f"Agenda - {data}")
    y = 750
    for cli in clientes:
        c.drawString(50, y, f"{cli[1]} - {cli[0]} ({cli[2]}) - Barbeiro: {cli[3]}")
        y -= 20
    c.save()
    output.seek(0)
    return send_file(output, download_name=f"agenda_{data}.pdf", as_attachment=True)

# -------------------------- FINANCEIRO ADMIN --------------------------
CATEGORIAS_FINANCEIRAS = ("Serviço", "Produto", "Despesa")


def _resolver_categoria_financeira(categoria, descricao, tipo_transacao):
    """Normaliza categoria para Serviço, Produto ou Despesa."""
    cat = (categoria or "").strip()
    if cat in CATEGORIAS_FINANCEIRAS:
        return cat
    if (tipo_transacao or "").strip() == "Despesa":
        return "Despesa"
    desc = (descricao or "").strip().lower()
    if desc.startswith("venda") or "produto" in desc[:24]:
        return "Produto"
    return "Serviço"


def _categoria_de_tipo_atendimento(tipo_feito):
    tipo = (tipo_feito or "").strip()
    if tipo in ("Venda de Produto", "Produto"):
        return "Produto"
    return "Serviço"


def _tipo_transacao_de_categoria(categoria):
    return "Despesa" if categoria == "Despesa" else "Receita"


def _financeiro_tem_coluna(cursor, coluna):
    """Verifica coluna em financeiro sem quebrar se a tabela ainda não migrou."""
    try:
        return _coluna_existe(cursor, "financeiro", coluna)
    except Exception:
        return False


def _formatar_transacao_financeira(row):
    """Converte linha SQL em dict para templates/exportação."""
    descricao = _valor_linha(row, 0) if row is not None else ""
    valor = float(_valor_linha(row, 1) or 0)
    tipo_transacao = _valor_linha(row, 2) or "Receita"
    profissional = _valor_linha(row, 3) or "Geral / Estabelecimento"
    data = _valor_linha(row, 4)
    categoria_raw = _valor_linha(row, 6, "categoria")
    categoria = _resolver_categoria_financeira(
        categoria_raw, descricao, tipo_transacao
    )
    return {
        "descricao": descricao,
        "valor": valor,
        "tipo_transacao": tipo_transacao,
        "profissional": profissional,
        "data": data,
        "profissional_id": _valor_linha(row, 5, "profissional_id"),
        "categoria": categoria,
    }


def _buscar_transacoes_financeiro(cursor, barbearia_id, profissional_id=None):
    """Lançamentos do tenant, opcionalmente filtrados por profissional."""
    cat_sel = (
        "f.categoria"
        if _financeiro_tem_coluna(cursor, "categoria")
        else "NULL AS categoria"
    )
    sql = f"""
        SELECT f.descricao, f.valor, f.tipo_transacao,
               COALESCE(u.nome, f.barbeiro, 'Geral / Estabelecimento') AS profissional,
               f.data, f.profissional_id, {cat_sel}
        FROM financeiro f
        LEFT JOIN usuarios u ON f.profissional_id = u.id
        WHERE f.barbearia_id = ?
    """
    params = [barbearia_id]
    if profissional_id:
        sql += " AND f.profissional_id = ?"
        params.append(int(profissional_id))
    sql += " ORDER BY f.data DESC"
    cursor.execute(sql, tuple(params))
    return [_formatar_transacao_financeira(r) for r in cursor.fetchall()]


def _calcular_saldo_financeiro(transacoes):
    saldo = 0.0
    for t in transacoes:
        if isinstance(t, dict):
            if t.get("tipo_transacao") == "Receita":
                saldo += float(t.get("valor", 0))
            else:
                saldo -= float(t.get("valor", 0))
        elif t[2] == "Receita":
            saldo += float(t[1])
        else:
            saldo -= float(t[1])
    return saldo


def _buscar_faturamento_profissionais(cursor, barbearia_id):
    """Total de receitas agrupado por profissional (legado / exportações)."""
    detalhado = _buscar_faturamento_detalhado_profissionais(cursor, barbearia_id)
    return [(r["nome"], r["total"]) for r in detalhado]


def _buscar_faturamento_detalhado_profissionais(cursor, barbearia_id):
    """Receitas por profissional com totais de Serviços e Produtos."""
    cat_sel = (
        "f.categoria"
        if _financeiro_tem_coluna(cursor, "categoria")
        else "NULL AS categoria"
    )
    cursor.execute(
        f"""
        SELECT u.id, u.nome, f.descricao, f.valor, f.tipo_transacao, {cat_sel}
        FROM financeiro f
        INNER JOIN usuarios u ON f.profissional_id = u.id
        WHERE f.barbearia_id = ? AND f.tipo_transacao = 'Receita'
        """,
        (barbearia_id,),
    )
    agregado = {}
    for row in cursor.fetchall() or []:
        uid = _valor_linha(row, 0, "id")
        nome = _valor_linha(row, 1, "nome")
        cat = _resolver_categoria_financeira(
            _valor_linha(row, 5, "categoria"),
            _valor_linha(row, 2, "descricao"),
            _valor_linha(row, 4, "tipo_transacao"),
        )
        valor = float(_valor_linha(row, 3, "valor") or 0)
        if uid not in agregado:
            agregado[uid] = {
                "id": uid,
                "nome": nome,
                "total_servicos": 0.0,
                "total_produtos": 0.0,
                "total": 0.0,
            }
        if cat == "Produto":
            agregado[uid]["total_produtos"] += valor
        else:
            agregado[uid]["total_servicos"] += valor
        agregado[uid]["total"] += valor
    return sorted(agregado.values(), key=lambda x: -x["total"])


def _inserir_receita_financeiro(
    cursor,
    descricao,
    valor,
    profissional_id,
    barbearia_id,
    agendamento_id=None,
    categoria="Serviço",
    nome_item=None,
):
    """Registra receita na tabela financeiro (uso pela agenda e lançamentos manuais)."""
    categoria = _resolver_categoria_financeira(categoria, descricao, "Receita")
    nome_prof, pid = _nome_profissional_lancamento(
        cursor, profissional_id, barbearia_id
    )
    cols = [
        "descricao",
        "valor",
        "tipo_transacao",
        "barbeiro",
        "profissional_id",
        "barbearia_id",
        "data",
    ]
    vals = [
        descricao,
        float(valor),
        "Receita",
        nome_prof,
        pid,
        int(barbearia_id),
        datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
    ]
    if _financeiro_tem_coluna(cursor, "categoria"):
        cols.append("categoria")
        vals.append(categoria)
    item_txt = (nome_item or "").strip()[:500] or None
    if categoria == "Produto" and _financeiro_tem_coluna(cursor, "produto"):
        cols.append("produto")
        vals.append(item_txt)
    elif _financeiro_tem_coluna(cursor, "servico"):
        cols.append("servico")
        vals.append(item_txt)
    if agendamento_id and _financeiro_tem_coluna(cursor, "agendamento_id"):
        cols.append("agendamento_id")
        vals.append(int(agendamento_id))
    placeholders = ", ".join("?" for _ in vals)
    colunas_sql = ", ".join(cols)
    cursor.execute(
        f"INSERT INTO financeiro ({colunas_sql}) VALUES ({placeholders})",
        tuple(vals),
    )


def _registrar_receita_agendamento(
    cursor,
    descricao,
    valor,
    profissional_id,
    barbearia_id,
    agendamento_id=None,
    substituir_existente=False,
    categoria="Serviço",
    nome_item=None,
):
    """
    Lança ou atualiza receita vinculada a um agendamento (evita duplicata ao concluir).
    """
    valor = float(valor or 0)
    categoria = _resolver_categoria_financeira(categoria, descricao, "Receita")
    if (
        substituir_existente
        and agendamento_id
        and _financeiro_tem_coluna(cursor, "agendamento_id")
    ):
        cursor.execute(
            """
            SELECT id FROM financeiro
            WHERE agendamento_id = ? AND barbearia_id = ?
            LIMIT 1
            """,
            (int(agendamento_id), int(barbearia_id)),
        )
        existente = cursor.fetchone()
        if existente:
            fin_id = _valor_linha(existente, 0, "id")
            nome_prof, pid = _nome_profissional_lancamento(
                cursor, profissional_id, barbearia_id
            )
            data_lanc = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
            if _financeiro_tem_coluna(cursor, "categoria"):
                cursor.execute(
                    """
                    UPDATE financeiro
                    SET descricao = ?, valor = ?, barbeiro = ?, profissional_id = ?,
                        tipo_transacao = 'Receita', categoria = ?, data = ?
                    WHERE id = ? AND barbearia_id = ?
                    """,
                    (
                        descricao,
                        valor,
                        nome_prof,
                        pid,
                        categoria,
                        data_lanc,
                        fin_id,
                        int(barbearia_id),
                    ),
                )
            else:
                cursor.execute(
                    """
                    UPDATE financeiro
                    SET descricao = ?, valor = ?, barbeiro = ?, profissional_id = ?,
                        tipo_transacao = 'Receita', data = ?
                    WHERE id = ? AND barbearia_id = ?
                    """,
                    (
                        descricao,
                        valor,
                        nome_prof,
                        pid,
                        data_lanc,
                        fin_id,
                        int(barbearia_id),
                    ),
                )
            return
    _inserir_receita_financeiro(
        cursor,
        descricao,
        valor,
        profissional_id,
        barbearia_id,
        agendamento_id=agendamento_id,
        categoria=categoria,
        nome_item=nome_item,
    )


def _montar_descricao_atendimento(tipo, detalhe, servico_agendado, nome_cliente):
    detalhe = (detalhe or "").strip()
    servico_agendado = (servico_agendado or "").strip()
    if tipo == "Serviço":
        base = detalhe or servico_agendado or "Atendimento"
        return f"Serviço: {base}"
    if tipo == "Venda de Produto":
        base = detalhe or "Produto"
        return f"Venda: {base}"
    if detalhe:
        return detalhe
    if servico_agendado:
        return f"Atendimento — {servico_agendado} ({nome_cliente})"
    return f"Atendimento — {nome_cliente}"


def _nome_profissional_lancamento(cursor, profissional_id, barbearia_id):
    if not profissional_id:
        return "Geral / Estabelecimento", None
    cursor.execute(
        """
        SELECT nome FROM usuarios
        WHERE id = ? AND barbearia_id = ?
          AND role IN ('barbeiro', 'profissional')
        """,
        (profissional_id, barbearia_id),
    )
    row = cursor.fetchone()
    if row:
        return row[0], profissional_id
    return "Geral / Estabelecimento", None


def _formatar_data_transacao(data):
    if hasattr(data, "strftime"):
        return data.strftime("%d/%m/%Y %H:%M")
    return str(data)[:16] if data else ""


def _gerar_excel_financeiro(transacoes, saldo, faturamento_detalhado, filtros=None):
    from openpyxl import Workbook

    filtros = filtros or {}
    wb = Workbook()
    ws = wb.active
    ws.title = "Financeiro"
    if filtros.get("data_ini") or filtros.get("data_fim"):
        ws.append(["Período", f"{filtros.get('data_ini', '')} a {filtros.get('data_fim', '')}"])
        ws.append([])
    ws.append(["Data", "Descrição", "Profissional", "Categoria", "Tags", "Tipo", "Valor"])
    for t in transacoes:
        if isinstance(t, dict):
            ws.append([
                _formatar_data_transacao(t["data"]),
                t["descricao"],
                t["profissional"],
                t["categoria"],
                t.get("tags", ""),
                t["tipo_transacao"],
                float(t["valor"]),
            ])
        else:
            ws.append([
                _formatar_data_transacao(t[4]),
                t[0],
                t[3],
                "",
                t[2],
                float(t[1]),
            ])
    ws.append([])
    ws.append(["", "", "", "", "Saldo em caixa", saldo])
    ws.append([])
    ws.append(["Faturamento por Profissional", ""])
    ws.append(["Profissional", "Serviços (R$)", "Produtos (R$)", "Total (R$)"])
    for row in faturamento_detalhado:
        ws.append([
            row["nome"],
            float(row["total_servicos"]),
            float(row["total_produtos"]),
            float(row["total"]),
        ])
    if not faturamento_detalhado:
        ws.append(["—", 0.0, 0.0, 0.0])
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output


def _gerar_pdf_financeiro(transacoes, saldo, titulo_estabelecimento, faturamento_detalhado):
    output = io.BytesIO()
    c = canvas.Canvas(output, pagesize=A4)
    largura, altura = A4
    y = altura - 2 * cm

    c.setFont("Helvetica-Bold", 16)
    c.drawString(2 * cm, y, titulo_estabelecimento)
    y -= 0.8 * cm
    c.setFont("Helvetica-Bold", 12)
    c.drawString(2 * cm, y, "Relatório Financeiro")
    y -= 0.6 * cm
    c.setFont("Helvetica", 11)
    c.drawString(2 * cm, y, f"Saldo em caixa: R$ {saldo:.2f}")
    y -= 1 * cm

    colunas = ["Data", "Descrição", "Profissional", "Categoria", "Valor"]
    xs = [2 * cm, 3.8 * cm, 8.5 * cm, 12.5 * cm, 16 * cm]
    c.setFont("Helvetica-Bold", 9)
    for i, col in enumerate(colunas):
        c.drawString(xs[i], y, col)
    y -= 0.5 * cm
    c.line(2 * cm, y, largura - 2 * cm, y)
    y -= 0.4 * cm

    c.setFont("Helvetica", 8)
    for t in transacoes:
        if y < 2 * cm:
            c.showPage()
            y = altura - 2 * cm
            c.setFont("Helvetica", 8)
        if isinstance(t, dict):
            linha = [
                _formatar_data_transacao(t["data"])[:10],
                (t["descricao"] or "")[:24],
                (t["profissional"] or "")[:16],
                (t["categoria"] or "")[:12],
                f"R$ {float(t['valor']):.2f}",
            ]
        else:
            linha = [
                _formatar_data_transacao(t[4])[:10],
                (t[0] or "")[:24],
                (t[3] or "")[:16],
                t[2] or "",
                f"R$ {float(t[1]):.2f}",
            ]
        for i, texto in enumerate(linha):
            c.drawString(xs[i], y, texto)
        y -= 0.45 * cm

    y -= 0.8 * cm
    if y < 4 * cm:
        c.showPage()
        y = altura - 2 * cm

    c.setFont("Helvetica-Bold", 11)
    c.drawString(2 * cm, y, "Faturamento por Profissional")
    y -= 0.6 * cm
    c.setFont("Helvetica-Bold", 8)
    c.drawString(2 * cm, y, "Profissional")
    c.drawString(7.5 * cm, y, "Serviços")
    c.drawString(11 * cm, y, "Produtos")
    c.drawString(14.5 * cm, y, "Total")
    y -= 0.4 * cm
    c.line(2 * cm, y, largura - 2 * cm, y)
    y -= 0.35 * cm

    c.setFont("Helvetica", 8)
    if faturamento_detalhado:
        for row in faturamento_detalhado:
            if y < 2 * cm:
                c.showPage()
                y = altura - 2 * cm
                c.setFont("Helvetica", 8)
            c.drawString(2 * cm, y, (row["nome"] or "")[:28])
            c.drawString(7.5 * cm, y, f"{float(row['total_servicos']):.2f}")
            c.drawString(11 * cm, y, f"{float(row['total_produtos']):.2f}")
            c.drawString(14.5 * cm, y, f"{float(row['total']):.2f}")
            y -= 0.4 * cm
    else:
        c.drawString(2 * cm, y, "Nenhuma receita vinculada a profissionais.")
        y -= 0.4 * cm

    c.save()
    output.seek(0)
    return output


@app.route("/admin_financeiro")
@requer_plano
def admin_financeiro():
    bloqueio = _exigir_admin()
    if bloqueio:
        return bloqueio

    barbearia_id = _barbearia_id_admin_obrigatorio()
    if not barbearia_id:
        return redirect(url_for("login"))

    filtros = fin.parse_filtros_request(request.args)
    filtro_profissional_id = filtros["profissional_id"]

    try:
        db_adapter.ensure_financeiro_schema()
        profissionais = fin.listar_profissionais(barbearia_id)

        if filtro_profissional_id and not fin.profissional_pertence_barbearia(
            filtro_profissional_id, barbearia_id
        ):
            filtro_profissional_id = None
            flash(_("Profissional inválido para este negócio."), "warning")

        comissoes = fin.carregar_comissoes(barbearia_id)
        transacoes = fin.buscar_transacoes_financeiro(
            barbearia_id,
            filtro_profissional_id,
            filtros["data_ini"],
            filtros["data_fim"],
        )
        kpis = fin.calcular_kpis(transacoes, comissoes)
        fechamento = fin.calcular_fechamento_caixa(
            transacoes, comissoes, filtro_profissional_id
        )
        faturamento_detalhado = fin.buscar_faturamento_detalhado_profissionais(
            barbearia_id,
            filtros["data_ini"],
            filtros["data_fim"],
        )
        grafico = fin.faturamento_diario_para_grafico(
            barbearia_id,
            filtros["data_ini"],
            filtros["data_fim"],
        )
        profissionais_por_nome = {p[1]: p[0] for p in profissionais}

        return render_template(
            "admin_financeiro.html",
            transacoes=transacoes,
            kpis=kpis,
            comissoes=comissoes,
            fechamento=fechamento,
            profissionais=profissionais,
            profissionais_por_nome=profissionais_por_nome,
            faturamento_detalhado=faturamento_detalhado,
            filtro_profissional_id=filtro_profissional_id,
            data_ini=filtros["data_ini"],
            data_fim=filtros["data_fim"],
            grafico_labels=grafico["labels"],
            grafico_valores=grafico["valores"],
            comissoes_json=comissoes,
        )
    except Exception:
        app.logger.exception("Erro no Financeiro")
        flash(
            _(
                "Não foi possível carregar o financeiro. "
                "Confira o terminal (Erro no Financeiro) e as migrações do banco."
            ),
            "danger",
        )
        return redirect(url_for("admin_agenda"))


@app.route("/admin_financeiro/comissoes", methods=["POST"])
@requer_plano
def admin_financeiro_salvar_comissoes():
    bloqueio = _exigir_admin()
    if bloqueio:
        return bloqueio
    barbearia_id = _barbearia_id_admin_obrigatorio()
    if not barbearia_id:
        return redirect(url_for("login"))
    try:
        pct_servico = float(
            str(request.form.get("pct_servico", fin.DEFAULT_PCT_SERVICO)).replace(",", ".")
        )
        pct_produto = float(
            str(request.form.get("pct_produto", fin.DEFAULT_PCT_PRODUTO)).replace(",", ".")
        )
    except (TypeError, ValueError):
        flash(_("Percentuais de comissão inválidos."), "danger")
        return redirect(url_for("admin_financeiro"))
    try:
        db_adapter.ensure_financeiro_schema()
        fin.salvar_comissoes(barbearia_id, pct_servico, pct_produto)
        flash(_("Comissões atualizadas com sucesso."), "success")
    except Exception:
        app.logger.exception("Erro ao salvar comissões")
        flash(_("Não foi possível salvar as comissões."), "danger")
    q = request.form.get("redirect_query", "")
    return redirect(url_for("admin_financeiro") + (f"?{q}" if q else ""))


@app.route("/lancar_transacao", methods=["POST"])
@requer_plano
def lancar_transacao():
    bloqueio = _exigir_admin()
    if bloqueio:
        return bloqueio

    barbearia_id = _barbearia_id_admin_obrigatorio()
    if not barbearia_id:
        return redirect(url_for("login"))

    descricao = request.form["descricao"]
    valor = request.form["valor"]
    categoria_raw = (request.form.get("categoria") or "Serviço").strip()
    profissional_id_raw = request.form.get("profissional_id", "").strip()
    profissional_id = int(profissional_id_raw) if profissional_id_raw else None

    if categoria_raw not in fin.CATEGORIAS_FINANCEIRAS:
        flash(_("Categoria inválida."), "error")
        return redirect(url_for("admin_financeiro"))
    categoria = categoria_raw
    tipo = fin.tipo_transacao_de_categoria(categoria)
    tags = fin.normalizar_tags(request.form.get("tags", ""))

    try:
        valor_num = float(str(valor).strip().replace(",", "."))
    except (TypeError, ValueError):
        flash(_("Valor inválido."), "error")
        return redirect(url_for("admin_financeiro"))

    try:
        db_adapter.ensure_financeiro_schema()
        nome_profissional, profissional_id = fin.nome_profissional_lancamento(
            profissional_id, barbearia_id
        )
        data_lanc = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        fin.inserir_lancamento_financeiro(
            barbearia_id,
            descricao,
            valor_num,
            tipo,
            categoria,
            nome_profissional,
            profissional_id,
            data_lanc,
            tags,
        )
    except Exception:
        app.logger.exception("Erro ao lançar transação financeira")
        flash(_("Não foi possível salvar o lançamento."), "danger")
        return redirect(url_for("admin_financeiro"))

    q = request.form.get("redirect_query", "").strip()
    return redirect(url_for("admin_financeiro") + (f"?{q}" if q else ""))


@app.route("/financeiro/exportar/excel")
@requer_plano
def financeiro_exportar_excel():
    bloqueio = _exigir_admin()
    if bloqueio:
        return bloqueio

    barbearia_id = _barbearia_id_admin_obrigatorio()
    if not barbearia_id:
        return redirect(url_for("login"))

    filtros = fin.parse_filtros_request(request.args)
    filtro_prof = filtros["profissional_id"]

    db_adapter.ensure_financeiro_schema()
    if filtro_prof and not fin.profissional_pertence_barbearia(
        filtro_prof, barbearia_id
    ):
        filtro_prof = None
    transacoes = fin.buscar_transacoes_financeiro(
        barbearia_id,
        filtro_prof,
        filtros["data_ini"],
        filtros["data_fim"],
    )
    kpis = fin.calcular_kpis(transacoes, fin.carregar_comissoes(barbearia_id))
    faturamento_detalhado = fin.buscar_faturamento_detalhado_profissionais(
        barbearia_id,
        filtros["data_ini"],
        filtros["data_fim"],
    )

    output = _gerar_excel_financeiro(
        transacoes, kpis["saldo_caixa"], faturamento_detalhado, filtros
    )
    ini = filtros.get("data_ini", "")[:10].replace("-", "")
    fim = filtros.get("data_fim", "")[:10].replace("-", "")
    nome = f"financeiro_{ini}_{fim}.xlsx" if ini and fim else f"financeiro_{datetime.today().strftime('%Y%m%d')}.xlsx"
    return send_file(
        output,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        as_attachment=True,
        download_name=nome,
    )


@app.route("/financeiro/exportar/pdf")
@requer_plano
def financeiro_exportar_pdf():
    bloqueio = _exigir_admin()
    if bloqueio:
        return bloqueio

    barbearia_id = _barbearia_id_admin_obrigatorio()
    if not barbearia_id:
        return redirect(url_for("login"))

    titulo = session.get("nome_barbearia", "AgendaSimples")
    filtros = fin.parse_filtros_request(request.args)
    filtro_prof = filtros["profissional_id"]

    db_adapter.ensure_financeiro_schema()
    if filtro_prof and not fin.profissional_pertence_barbearia(
        filtro_prof, barbearia_id
    ):
        filtro_prof = None
    transacoes = fin.buscar_transacoes_financeiro(
        barbearia_id,
        filtro_prof,
        filtros["data_ini"],
        filtros["data_fim"],
    )
    kpis = fin.calcular_kpis(transacoes, fin.carregar_comissoes(barbearia_id))
    faturamento_detalhado = fin.buscar_faturamento_detalhado_profissionais(
        barbearia_id,
        filtros["data_ini"],
        filtros["data_fim"],
    )

    output = _gerar_pdf_financeiro(
        transacoes, kpis["saldo_caixa"], titulo, faturamento_detalhado
    )
    nome = f"financeiro_{datetime.today().strftime('%Y%m%d')}.pdf"
    return send_file(
        output,
        mimetype="application/pdf",
        as_attachment=True,
        download_name=nome,
    )

# -------------------------- DINÂMICA DA GALERIA --------------------------
@app.route("/admin/galeria", methods=["GET", "POST"])
@requer_plano
def admin_galeria():
    bloqueio = _exigir_admin()
    if bloqueio:
        return bloqueio

    barbearia_id = _barbearia_id_admin_obrigatorio()
    if not barbearia_id:
        return redirect(url_for("login"))

    conn, cursor = _open_db()

    if request.method == "POST":
        categoria = request.form.get("categoria")
        file = request.files.get("foto")

        if file and file.filename != '':
            upload_dir = 'static/uploads/galeria'
            if not os.path.exists(upload_dir):
                os.makedirs(upload_dir)
                
            extensao = file.filename.rsplit('.', 1)[1].lower() if '.' in file.filename else 'jpg'
            timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
            filename = f"foto_{categoria}_{timestamp}.{extensao}"
            
            file.save(os.path.join(upload_dir, filename))

            cursor.execute(
                "INSERT INTO tb_galeria (barbearia_id, categoria, caminho_foto) VALUES (?, ?, ?)",
                (barbearia_id, categoria, filename),
            )
            safe_commit(conn)

    todas_fotos = db_adapter.execute_query(
        """
        SELECT id, categoria, caminho_foto FROM tb_galeria
        WHERE barbearia_id = ? ORDER BY id DESC
        """,
        (barbearia_id,),
        conn=conn,
    )
    safe_close(conn)
    
    galeria = {"corte": [], "corte_barba": [], "sobrancelha": [], "outros": []}
    for f in todas_fotos or []:
        cat = db_adapter.row_get(f, "categoria", index=1)
        if cat in galeria:
            galeria[cat].append({
                "id": db_adapter.row_get(f, "id", index=0),
                "foto": db_adapter.row_get(f, "caminho_foto", index=2),
            })
            
    return render_template("admin_galeria.html", galeria=galeria)

@app.route("/eliminar_foto/<int:foto_id>", methods=["POST", "GET"])
@requer_plano
def eliminar_foto(foto_id):
    bloqueio = _exigir_admin()
    if bloqueio:
        return bloqueio

    barbearia_id = _barbearia_id_admin_obrigatorio()
    if not barbearia_id:
        return redirect(url_for("login"))

    conn, cursor = _open_db()

    foto = db_adapter.execute_query(
        "SELECT caminho_foto FROM tb_galeria WHERE id = ? AND barbearia_id = ?",
        (foto_id, barbearia_id),
        fetch="one",
        conn=conn,
    )
    
    if foto:
        nome_arquivo = db_adapter.row_get(foto, "caminho_foto", index=0)
        caminho_completo = os.path.join('static/uploads/galeria', nome_arquivo)
        
        if os.path.exists(caminho_completo):
            try:
                os.remove(caminho_completo)
            except Exception as e:
                print(f"Erro ao deletar arquivo físico: {e}")
            
        cursor.execute(
            "DELETE FROM tb_galeria WHERE id = ? AND barbearia_id = ?",
            (foto_id, barbearia_id),
        )
        safe_commit(conn)
        
    safe_close(conn)
    return redirect(url_for("admin_galeria"))

# -------------------------- CONFIGURAÇÕES DO ADMIN (COM FOTO DE CAPA) --------------------------
@app.route("/admin/configuracoes", methods=["GET", "POST"])
@requer_plano
def admin_configuracoes():
    bloqueio = _exigir_admin()
    if bloqueio:
        return bloqueio
        
    conn, cursor = _open_db()
    
    barbearia_id_alvo = _barbearia_id_admin_obrigatorio()
    if not barbearia_id_alvo:
        safe_close(conn)
        return redirect(url_for("login"))

    if request.method == "POST":
        novo_nome = request.form.get("nome_barbearia")
        file_capa = request.files.get("foto_capa")
        file_logotipo = request.files.get("logotipo")
        titulo1 = request.form.get("titulo_catalogo1", "").strip()
        titulo2 = request.form.get("titulo_catalogo2", "").strip()
        titulo3 = request.form.get("titulo_catalogo3", "").strip()
        titulo4 = request.form.get("titulo_catalogo4", "").strip()
        texto_marcar = request.form.get("texto_marcar_direito", "").strip()
        file_fundo_marcar = request.files.get("foto_fundo_marcar")
        link_instagram = _normalizar_link_url(request.form.get("link_instagram", ""))
        link_facebook = _normalizar_link_url(request.form.get("link_facebook", ""))
        link_whatsapp = _normalizar_link_whatsapp(request.form.get("link_whatsapp", ""))

        # 1. Salva nome e títulos dos catálogos da Home
        if novo_nome:
            cursor.execute("UPDATE barbearias SET nome = ? WHERE id = ?", (novo_nome, barbearia_id_alvo))
            session['nome_barbearia'] = novo_nome

        try:
            cursor.execute(
                """
                UPDATE barbearias SET
                    titulo_catalogo1 = ?,
                    titulo_catalogo2 = ?,
                    titulo_catalogo3 = ?,
                    titulo_catalogo4 = ?
                WHERE id = ?
                """,
                (
                    titulo1 or DEFAULT_TITULOS_CATALOGO[0],
                    titulo2 or DEFAULT_TITULOS_CATALOGO[1],
                    titulo3 or DEFAULT_TITULOS_CATALOGO[2],
                    titulo4 or DEFAULT_TITULOS_CATALOGO[3],
                    barbearia_id_alvo,
                ),
            )
        except DbError:
            pass

        _processar_upload_capas_catalogo(request, cursor, barbearia_id_alvo)

        try:
            cursor.execute(
                """
                UPDATE barbearias SET
                    texto_marcar_direito = ?,
                    link_instagram = ?,
                    link_facebook = ?,
                    link_whatsapp = ?
                WHERE id = ?
                """,
                (
                    texto_marcar or novo_nome or "AgendaSimples",
                    link_instagram,
                    link_facebook,
                    link_whatsapp,
                    barbearia_id_alvo,
                ),
            )
        except DbError:
            try:
                cursor.execute(
                    "UPDATE barbearias SET texto_marcar_direito = ? WHERE id = ?",
                    (texto_marcar or novo_nome or "AgendaSimples", barbearia_id_alvo),
                )
            except DbError:
                pass

        if file_fundo_marcar and file_fundo_marcar.filename:
            upload_banner_dir = "static/uploads/banners"
            os.makedirs(upload_banner_dir, exist_ok=True)
            ext = (
                file_fundo_marcar.filename.rsplit(".", 1)[1].lower()
                if "." in file_fundo_marcar.filename
                else "jpg"
            )
            banner_filename = f"fundo_direito_{barbearia_id_alvo}.{ext}"
            file_fundo_marcar.save(os.path.join(upload_banner_dir, banner_filename))
            try:
                cursor.execute(
                    """
                    UPDATE barbearias SET foto_fundo_direito_url = ? WHERE id = ?
                    """,
                    (banner_filename, barbearia_id_alvo),
                )
            except DbError:
                pass

        # 2. Upload do logotipo (Home)
        if file_logotipo and file_logotipo.filename:
            upload_logo_dir = "static/uploads/logo"
            os.makedirs(upload_logo_dir, exist_ok=True)
            ext = (
                file_logotipo.filename.rsplit(".", 1)[1].lower()
                if "." in file_logotipo.filename
                else "png"
            )
            logo_filename = f"logo_barbearia_{barbearia_id_alvo}.{ext}"
            file_logotipo.save(os.path.join(upload_logo_dir, logo_filename))
            try:
                cursor.execute(
                    "UPDATE barbearias SET logotipo_url = ? WHERE id = ?",
                    (logo_filename, barbearia_id_alvo),
                )
            except DbError:
                pass

        # 3. Processa o upload da Foto de Capa se o usuário escolheu um arquivo
        if file_capa and file_capa.filename != '':
            upload_dir = 'static/uploads/capa'
            if not os.path.exists(upload_dir):
                os.makedirs(upload_dir)
                
            extensao = file_capa.filename.rsplit('.', 1)[1].lower() if '.' in file_capa.filename else 'jpg'
            filename = f"capa_barbearia_{barbearia_id_alvo}.{extensao}"
            
            # Salva o arquivo físico na pasta
            file_capa.save(os.path.join(upload_dir, filename))
            
            # Verifica se já existe o registro 'foto_capa' na tb_configuracoes
            cursor.execute("SELECT COUNT(*) FROM tb_configuracoes WHERE chave = 'foto_capa'")
            existe_config = cursor.fetchone()[0]
            
            if existe_config > 0:
                cursor.execute("UPDATE tb_configuracoes SET valor = ? WHERE chave = 'foto_capa'", (filename,))
            else:
                cursor.execute("INSERT INTO tb_configuracoes (chave, valor) VALUES ('foto_capa', ?)", (filename,))
        
        conn.commit()
        conn.close()
        return redirect(url_for("admin_configuracoes"))
        
    # --- GERENCIAMENTO DO SINAL DO GET (CARREGAMENTO DA TELA) ---
    # Busca o nome atual da barbearia
    configs = _carregar_configs_home(cursor, barbearia_id_alvo)

    cursor.execute("SELECT valor FROM tb_configuracoes WHERE chave = 'foto_capa'")
    config_foto = cursor.fetchone()
    foto_capa = config_foto[0] if config_foto else None

    conn.close()
    return render_template(
        "admin_configuracoes.html",
        nome_atual=configs["nome_negocio"],
        foto_capa=foto_capa,
        logotipo_url=configs["logotipo_url"],
        titulo_catalogo1=configs["titulo_catalogo1"],
        titulo_catalogo2=configs["titulo_catalogo2"],
        titulo_catalogo3=configs["titulo_catalogo3"],
        titulo_catalogo4=configs["titulo_catalogo4"],
        texto_marcar_direito=configs.get("texto_marcar_direito") or configs["nome_negocio"],
        foto_fundo_direito_url=configs.get("foto_fundo_direito_url"),
        capa_catalogo1=configs.get("capa_catalogo1"),
        capa_catalogo2=configs.get("capa_catalogo2"),
        capa_catalogo3=configs.get("capa_catalogo3"),
        capa_catalogo4=configs.get("capa_catalogo4"),
        link_instagram=configs.get("link_instagram") or "",
        link_facebook=configs.get("link_facebook") or "",
        link_whatsapp=configs.get("link_whatsapp") or "",
    )

@app.route("/admin/clientes")
@requer_plano
def admin_clientes():
    return render_template("admin_clientes.html")

@app.route("/admin/assinatura")
@requer_plano
def admin_assinatura():
    """Status do plano (rota liberada pelo decorator quando trial/plano ativo)."""
    from subscriptions import obter_assinatura, assinatura_permite_acesso

    barbearia_id = session.get("barbearia_id") or session.get("user_id")
    conn, cursor = _open_db()
    assinatura = obter_assinatura(cursor, barbearia_id)
    safe_close(conn)
    return render_template(
        "admin_assinatura.html",
        assinatura=assinatura,
        acesso_ok=assinatura_permite_acesso(assinatura),
        trial_days=cfg.TRIAL_DAYS,
    )

# -------------------------- RODAR --------------------------
if __name__ == "__main__":
    local_ip = socket.gethostbyname(socket.gethostname())
    print(f"Acesse em: http://{local_ip}:5000")
    app.run(host="0.0.0.0", port=5000, debug=True)